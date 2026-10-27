"""Certification campaign reports."""

import zipfile
from datetime import datetime, timezone

import pytest

from conftest import FakeApi, FakeIsc
from isc_reporter import certification_reports as cr
from isc_reporter import cli

NOW = datetime(2026, 10, 26, tzinfo=timezone.utc)


def campaign(cid, name, status, deadline, total, done, **extra):
    return {"id": cid, "name": name, "type": "MANAGER", "status": status, "created": "2026-09-01T08:00:00Z",
            "deadline": deadline, "totalCertifications": total, "completedCertifications": done,
            "autoRevokeAllowed": True, **extra}


def cert(cid, campaign_id, campaign_name, reviewer, made, total, due, completed=False, signed=None,
         approved=(0, 0), revoked=(0, 0)):
    """approved/revoked: (entitlements, roles)."""
    return {"id": cid, "campaign": {"id": campaign_id, "name": campaign_name},
            "reviewer": {"name": reviewer, "email": f"{reviewer.lower()}@example.com"},
            "identitiesCompleted": 1, "identitiesTotal": 2, "decisionsMade": made, "decisionsTotal": total,
            "due": due, "signed": signed, "completed": completed, "phase": "ACTIVE",
            "decisionSummary": {
                "entitlementsApproved": approved[0], "rolesApproved": approved[1],
                "entitlementsRevoked": revoked[0], "rolesRevoked": revoked[1],
                "entitlementDecisionsMade": approved[0] + revoked[0],
                "roleDecisionsMade": approved[1] + revoked[1],
                "entitlementDecisionsTotal": total - 2, "roleDecisionsTotal": 2}}


def item(identity, access_type, name, decision, campaign_name="Q3 Manager", reviewer="Ann", **extra):
    return {"identitySummary": {"name": identity}, "decision": decision, "comments": extra.get("comments"),
            "accessSummary": {"access": {"type": access_type, "name": name},
                              "entitlement": {"name": name, "sourceName": "AD", "privileged": True}
                              if access_type == "ENTITLEMENT" else None},
            "certification": {"campaign": campaign_name, "reviewer": reviewer}}


@pytest.fixture
def data():
    return cr.CertData(
        campaigns=[campaign("c1", "Q3 Manager", "ACTIVE", "2026-10-20T00:00:00Z", 2, 1,
                            filter={"id": "f1", "type": "CAMPAIGN_FILTER"}),
                   campaign("c2", "Q2 Manager", "COMPLETED", "2026-07-01T00:00:00Z", 1, 1),
                   campaign("c3", "Q4 AD", "STAGED", "2026-12-31T00:00:00Z", 0, 0)],
        certifications=[
            cert("x1", "c1", "Q3 Manager", "Ann", 6, 10, "2026-10-20T00:00:00Z",
                 approved=(3, 1), revoked=(2, 0)),
            cert("x2", "c1", "Q3 Manager", "Bob", 4, 4, "2026-10-20T00:00:00Z", completed=True,
                 signed="2026-10-18T10:00:00Z", approved=(2, 1), revoked=(0, 1)),
            cert("x3", "c2", "Q2 Manager", "Ann", 5, 5, "2026-07-01T00:00:00Z", completed=True,
                 signed="2026-06-30T10:00:00Z", approved=(3, 2))],
        items=[item("jdoe", "ENTITLEMENT", "Domain Admins", "REVOKE", comments="left the team"),
               item("jdoe", "ROLE", "Finance Analyst", "APPROVE"),
               item("jdoe", "ENTITLEMENT", "VPN", None),
               item("asmith", "ACCESS_PROFILE", "Mail", "APPROVE")],
        filters=[{"id": "f1", "name": "Only AD", "mode": "INCLUSION",
                  "criteriaList": [{"type": "SOURCE", "operation": "EQUALS", "value": "AD"}]}],
        with_items=True, now=NOW, origin="https://tenant", scope="all statuses")


def test_campaign_rows(data):
    rows = {r["Campaign"]: r for r in cr.campaign_rows(data)}
    q3 = rows["Q3 Manager"]
    assert (q3["Progress"], q3["Decisions made"], q3["Decisions"]) == (0.5, 10, 14)
    assert (q3["Approved"], q3["Revoked"], q3["Revocation rate"]) == (7, 3, 0.3)
    assert (q3["Days left"], q3["Overdue"], q3["Filter"]) == (-6, "yes", "Only AD")
    assert rows["Q2 Manager"]["Days left"] is None and rows["Q2 Manager"]["Overdue"] == "no"
    assert rows["Q4 AD"]["Progress"] is None and rows["Q4 AD"]["Approved"] is None


def test_certification_and_reviewer_rows(data):
    certs = cr.certification_rows(data)
    assert [(c["Reviewer"], c["Pending"], c["Overdue"]) for c in certs] == [
        ("Ann", 4, "yes"), ("Bob", 0, "no"), ("Ann", 0, "no")]
    reviewers = cr.reviewer_rows(data)
    ann = reviewers[0]
    assert (ann["Reviewer"], ann["Certifications"], ann["Signed off"], ann["Pending"]) == ("Ann", 2, 1, 4)
    assert (ann["Approved"], ann["Revoked"], ann["Overdue certifications"]) == (9, 2, 1)
    assert ann["Progress"] == pytest.approx(11 / 15)


def test_items_identities_and_filters(data):
    items = cr.item_rows(data)
    assert [(i["Access type"], i["Decision"]) for i in items] == [
        ("Entitlement", "Revoked"), ("Role", "Approved"), ("Entitlement", "Pending"),
        ("Access Profile", "Approved")]
    first = items[0]
    assert (first["Source"], first["Privileged"], first["Comments"]) == ("AD", "yes", "left the team")
    identities = {r["Identity"]: r for r in cr.identity_rows(data)}
    jdoe = identities["jdoe"]
    assert (jdoe["Approved"], jdoe["Revoked"], jdoe["Pending"]) == (1, 1, 1)
    assert identities["jdoe"]["Revoked access"] == "Domain Admins"
    [only_ad] = cr.filter_rows(data)
    assert only_ad["Criteria"] == "SOURCE EQUALS AD"
    assert only_ad["Used by campaigns in this report"] == "Q3 Manager"


def test_kpis_and_charts(data):
    tiles = {label: value for label, value, _ in cr.kpis(data)}
    assert tiles["Campaigns"] == "3  (1 active)"
    assert tiles["Certifications completed"] == "2 / 3  (67%)"
    assert tiles["Decisions made"] == "15 / 19  (79%)"
    assert tiles["Revoked"] == "3  (20% of decisions)"
    assert tiles["Overdue certifications"] == "1"
    charts = {c.title: c for c in cr.charts(data)}
    assert charts["Campaigns by status"].categories == ["STAGED", "ACTIVE", "COMPLETED"]
    assert charts["Decisions"].series[0][1] == [12, 3, 4]
    by_type = charts["Decisions by access type"]
    assert by_type.categories == ["Entitlements", "Roles"]
    assert [s[1] for s in by_type.series] == [[8, 4], [2, 1], [3, 1]]
    assert charts["Revocations by access type"].categories == ["Entitlement"]


def test_excel_and_pdf(data, tmp_path):
    from openpyxl import load_workbook

    path = cr.write_xlsx(data, str(tmp_path / "r.xlsx"), "Campaigns")
    book = load_workbook(path)
    assert book.sheetnames == ["Dashboard", "Campaigns", "Certifications", "Reviewers", "Identities",
                               "Decisions", "Filters", "Chart data", "Read me"]
    assert book["Campaigns"]["A2"].value == "Q3 Manager"
    package = zipfile.ZipFile(path)
    names = package.namelist()
    assert sum(n.startswith("xl/charts/chart") for n in names) == len(cr.charts(data))
    assert "doughnut" in package.read("xl/charts/chart1.xml").decode().lower()
    assert "xl/sharedStrings.xml" in names
    pdf = cr.write_pdf(data, str(tmp_path / "r.pdf"), "Campaigns")
    assert open(pdf, "rb").read(4) == b"%PDF"


def test_empty_campaign_set_renders(tmp_path):
    empty = cr.CertData(campaigns=[campaign("c9", "New", "STAGED", None, 0, 0)], now=NOW)
    cr.write_xlsx(empty, str(tmp_path / "e.xlsx"), "Empty")
    cr.write_pdf(empty, str(tmp_path / "e.pdf"), "Empty")


def test_fetch_and_cli_with_sdk_models(tmp_path, monkeypatch):
    from sailpoint.certification_campaign_filters.models import ListCampaignFiltersV1200Response
    from sailpoint.certification_campaigns.models import SlimCampaign
    from sailpoint.certification_summaries.models import IdentityCertDecisionSummary
    from sailpoint.certifications.models import AccessReviewItem, IdentityCertificationDto

    slim = SlimCampaign.from_dict({"id": "c1", "name": "Q3 Manager", "description": "d", "type": "MANAGER",
                                   "status": "ACTIVE", "created": "2026-09-01T08:00:00.000Z",
                                   "deadline": "2026-11-30T00:00:00.000Z", "totalCertifications": 1,
                                   "completedCertifications": 0})
    dto = IdentityCertificationDto.from_dict({
        "id": "x1", "name": "Ann's certification", "completed": False, "identitiesCompleted": 0,
        "identitiesTotal": 1, "decisionsMade": 1, "decisionsTotal": 2, "created": "2026-09-01T08:00:00.000Z",
        "due": "2026-11-30T00:00:00.000Z", "phase": "ACTIVE", "hasErrors": False,
        "campaign": {"id": "c1", "name": "Q3 Manager", "type": "CAMPAIGN", "campaignType": "MANAGER",
                     "correlatedStatus": "CORRELATED", "mandatoryCommentRequirement": "NO_DECISIONS"},
        "reviewer": {"id": "r", "name": "Ann", "email": "ann@example.com", "type": "IDENTITY"}})
    summary = IdentityCertDecisionSummary.from_dict({"entitlementsApproved": 1, "entitlementDecisionsMade": 1,
                                                     "entitlementDecisionsTotal": 2})
    review = AccessReviewItem.from_dict({"id": "i1", "completed": True, "decision": "APPROVE",
                                         "identitySummary": {"id": "s", "name": "jdoe", "identityId": "j"},
                                         "accessSummary": {"access": {"type": "ENTITLEMENT", "id": "e",
                                                                      "name": "VPN"}}})
    filters = ListCampaignFiltersV1200Response.from_dict({"items": [], "count": 0})

    def first_page(items):
        return lambda **kw: items if not kw.get("offset") else []

    isc = FakeIsc(
        campaigns=FakeApi(get_active_campaigns_v1=first_page([slim])),
        certifications=FakeApi(list_identity_certifications_v1=first_page([dto]),
                               list_identity_access_review_items_v1=first_page([review])),
        certification_summaries=FakeApi(get_identity_decision_summary_v1=lambda id: summary),
        campaign_filters=FakeApi(list_campaign_filters_v1=lambda **kw: filters))
    monkeypatch.setattr("isc_reporter.client.connect", lambda config=None: isc)
    out = tmp_path / "report"
    assert cli.main(["campaigns", "--status", "active", "--items", "-o", str(out),
                     "--log-dir", ""]) == 0
    assert (tmp_path / "report.xlsx").exists() and (tmp_path / "report.pdf").exists()
    assert isc.campaigns.called("get_active_campaigns_v1")[0]["filters"] == 'status in ("ACTIVE")'
    assert isc.certifications.called("list_identity_certifications_v1")[0]["filters"] == 'campaign.id eq "c1"'
    assert cli.main(["campaigns", "--campaign", "Nope", "-o", str(out), "--log-dir", ""]) == 2


def test_model_json_keeps_read_only_fields():
    """The SDK's to_dict drops id/status/created; reports need them."""
    from sailpoint.certification_campaigns.models import SlimCampaign

    from isc_reporter.api import model_json

    slim = SlimCampaign.from_dict({"id": "c1", "name": "n", "description": "d", "type": "MANAGER",
                                   "status": "ACTIVE", "created": "2026-09-01T08:00:00Z",
                                   "totalCertifications": 4, "completedCertifications": 1})
    data = model_json(slim)
    assert (data["id"], data["status"], data["totalCertifications"]) == ("c1", "ACTIVE", 4)
    assert data["created"].startswith("2026-09-01T08:00:00")
