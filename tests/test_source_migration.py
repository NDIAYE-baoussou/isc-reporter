"""Impact analysis of replacing the HR source."""

import json
import os

import pytest

from isc_reporter import cli
from isc_reporter.source_migration import NEW_SOURCE_ID, Source, analyse


def account(attribute, source="HR"):
    return {"type": "accountAttribute",
            "attributes": {"sourceName": source, "attributeName": attribute, "sourceId": source.lower() * 4}}


def ref(name, **attrs):
    return {"type": "reference", "attributes": {"id": name, **attrs}}


def mapping(name, definition):
    return {"identityAttributeName": name, "transformDefinition": definition}


def profile(name, source, *mappings, priority=10):
    return {"id": name.lower(), "name": name, "priority": priority, "identityCount": 5,
            "authoritativeSource": {"type": "SOURCE", "id": source.lower() * 4, "name": source},
            "identityAttributeConfig": {"enabled": True, "attributeTransforms": list(mappings)}}


PROFILES = [
    profile("IP_HR", "HR",
            mapping("uid", account("id")),
            mapping("email", account("mail", "AD")),
            mapping("endDate", ref("Format date", input=account("endDate"))),
            mapping("cloudLifecycleState", ref("LCS")),
            mapping("manager", account("manager"))),
    profile("IP_Contractors", "Contractors",
            mapping("uid", account("login", "Contractors")),
            mapping("department", account("department")), priority=20),
    profile("IP_Other", "Other", mapping("uid", account("id", "Other")), priority=30),
]
TRANSFORMS = {
    "Format date": {"id": "t1", "name": "Format date", "type": "dateFormat",
                    "attributes": {"outputFormat": "ISO8601"}},
    "HR status": {"id": "t2", "name": "HR status", "type": "lookup",
                  "attributes": {"table": {"A": "active", "default": "inactive"},
                                 "input": account("status")}},
    "LCS": {"id": "t3", "name": "LCS", "type": "firstValid",
            "attributes": {"values": [ref("HR status"), {"type": "static", "attributes": {"value": "x"}}]}},
    "Unused HR": {"id": "t4", "name": "Unused HR", "type": "upper", "attributes": {"input": account("x")}},
}
NEW = Source("Workday", "w" * 32, ["employeeId", "endDate", "manager", "status"])
OWNER = {"type": "IDENTITY", "id": "o" * 32, "name": "owner"}


def leaf(kind, attribute, value, source=None):
    key = {"type": kind, "property": f"attribute.{attribute}", **({"sourceId": source} if source else {})}
    return {"operation": "EQUALS", "key": key, "stringValue": value}


def role(name, criteria=None, enabled=True, **extra):
    membership = {"type": "STANDARD", "criteria": criteria} if criteria else None
    return {"id": name.lower(), "name": name, "enabled": enabled, "membership": membership, **extra}


HR_ID = "hr" * 4                                    # the id profile() gives to the HR source


def finance(source_id=HR_ID):
    return role("Finance", {"operation": "OR", "children": [
        {"operation": "AND", "children": [leaf("IDENTITY", "department", "Finance"),
                                          leaf("IDENTITY", "location", "Paris")]},
        leaf("ACCOUNT", "status", "A", source_id)]})


LEAVERS = {"id": "d1", "name": "Leavers", "parentId": "managers", "membership": {
    "type": "STANDARD", "criteria": leaf("IDENTITY", "cloudLifecycleState", "inactive")}}
ROLES = [
    finance(),
    role("VPN", leaf("IDENTITY", "location", "Paris")),                          # attribute not impacted
    role("Old HR group", leaf("ENTITLEMENT", "groups", "Staff", HR_ID), enabled=False),
    role("Listed", membership={"type": "IDENTITY_LIST",
                               "identities": [{"type": "IDENTITY", "id": "i" * 32}]}),
    role("Managers", dimensional=True, dimensions=[LEAVERS]),
]
SEGMENTS = [{"id": "s1", "name": "Finance visibility", "active": True, "visibilityCriteria": {"expression": {
    "operator": "AND", "children": [
        {"operator": "EQUALS", "attribute": "department", "value": {"type": "STRING", "value": "Finance"}},
        {"operator": "EQUALS", "attribute": "location", "value": {"type": "STRING", "value": "Paris"}}]}}}]


def assignment(identity, account):
    return {"property": identity, "value": account, "operation": "EQ"}


SOURCES = [
    {"id": "adad", "name": "AD", "correlationConfig": {"attributeAssignments": [
        assignment("uid", "sAMAccountName"), assignment("email", "mail")]}},
    {"id": HR_ID, "name": "HR", "correlationConfig": {"attributeAssignments": [assignment("uid", "id")]},
     "managerCorrelationMapping": {"accountAttributeName": "managerId", "identityAttributeName": "uid"}},
    {"id": "w" * 32, "name": "Workday",                                         # the new source: left alone
     "correlationConfig": {"attributeAssignments": [assignment("uid", "employeeId")]}},
]


@pytest.fixture
def plan():
    return analyse(PROFILES, TRANSFORMS, Source("HR"), NEW, {"id": "employeeId"})


def test_finds_impacted_profiles_and_transform_chain(plan):
    assert [(p["profile"]["name"], p["authoritative"], p["touched"]) for p in plan.profiles] == [
        ("IP_HR", True, ["uid", "endDate", "cloudLifecycleState", "manager"]),
        ("IP_Contractors", False, ["department"])]
    assert {n: i["impact"] for n, i in plan.transforms.items()} == {
        "HR status": "reads the old source", "Unused HR": "reads the old source",
        "LCS": "calls an impacted transform"}
    assert plan.transforms["LCS"]["used_by"] == ["IP_HR: cloudLifecycleState"]


def test_attribute_mapping_statuses(plan):
    status = {a: (i["new_attribute"], i["status"]) for a, i in plan.attributes.items()}
    assert status["id"] == ("employeeId", "from the mapping file")
    assert status["endDate"] == ("endDate", "same name in the new schema")
    assert status["status"] == ("status", "same name in the new schema")
    assert status["department"] == ("", "TO DEFINE")
    assert status["x"] == ("", "TO DEFINE")
    assert plan.attributes["department"]["used_by"] == ["IP_Contractors: department"]


def test_before_after_and_checks(plan):
    rows = {(m["profile"], m["identity_attribute"]): m for m in plan.mappings}
    assert rows[("IP_HR", "email")]["change"] == "unchanged"
    end = rows[("IP_HR", "endDate")]
    assert (end["before"], end["after"]) == ("Format date(HR.endDate)", "Format date(Workday.endDate)")
    assert rows[("IP_HR", "cloudLifecycleState")]["after"] == "LCS - Workday()"
    assert rows[("IP_Contractors", "department")]["change"] == "to complete"
    checks = {(c[1], c[2]) for c in plan.checks}
    assert ("department", "New attribute to define") in checks
    assert ("Unused HR", "Transform not used by a profile") in checks
    assert ("IP_HR", "Priority to choose") in checks


def test_drafts(plan):
    new_profile = plan.drafts["identity-profiles/IP_HR - Workday.json"]
    assert "id" not in new_profile and "identityCount" not in new_profile
    assert new_profile["authoritativeSource"] == {"type": "SOURCE", "id": "w" * 32, "name": "Workday"}
    uid = new_profile["identityAttributeConfig"]["attributeTransforms"][0]["transformDefinition"]
    assert uid["attributes"] == {"sourceName": "Workday", "attributeName": "employeeId", "sourceId": "w" * 32}
    # the original profile and transforms are left untouched
    original = PROFILES[0]["identityAttributeConfig"]["attributeTransforms"][0]
    assert original["transformDefinition"] == account("id")
    status = plan.drafts["transforms/HR status - Workday.json"]
    assert "id" not in status and status["attributes"]["input"]["attributes"]["sourceName"] == "Workday"
    lcs = plan.drafts["transforms/LCS - Workday.json"]
    assert lcs["attributes"]["values"][0]["attributes"]["id"] == "HR status - Workday"
    contractors = plan.drafts["identity-profiles/IP_Contractors (updated).json"]
    assert contractors["id"] == "ip_contractors"
    department = contractors["identityAttributeConfig"]["attributeTransforms"][1]["transformDefinition"]
    assert department["attributes"]["attributeName"] == "TODO_department"


def test_in_place_keeps_transform_names_and_ids():
    plan = analyse(PROFILES, TRANSFORMS, Source("HR"), Source("Workday"), in_place=True)
    lcs = plan.drafts["transforms/LCS.json"]
    assert lcs["id"] == "t3" and lcs["attributes"]["values"][0]["attributes"]["id"] == "HR status"
    status = plan.drafts["transforms/HR status.json"]["attributes"]["input"]["attributes"]
    assert status["sourceId"] == NEW_SOURCE_ID
    # without a schema for the new source, same names are assumed
    assert plan.attributes["department"]["status"] == "same name assumed (new schema unknown)"


def test_nothing_found_is_reported():
    plan = analyse(PROFILES, TRANSFORMS, Source("Nope"), NEW)
    assert not plan.profiles and ("", "Nothing found") in {(c[1], c[2]) for c in plan.checks}


def test_cli_offline(tmp_path):
    exports = tmp_path / "exports"
    exports.mkdir()
    (exports / "profiles.json").write_text(json.dumps(PROFILES), encoding="utf-8")
    (exports / "transforms.json").write_text(json.dumps(list(TRANSFORMS.values())), encoding="utf-8")
    attribute_map = tmp_path / "map.csv"
    attribute_map.write_text("old_attribute;new_attribute\nid;employeeId\ndepartment;\n", encoding="utf-8")
    out = tmp_path / "out"
    assert cli.main(["migration", "--old-source", "HR", "--new-source", "Workday",
                     "--from", str(exports), "--attribute-map", str(attribute_map), "-o", str(out),
                     "--log-dir", ""]) == 0
    assert sorted(os.listdir(out)) == ["attribute-mapping.csv", "identity-profiles",
                                       "source-migration.pdf", "source-migration.xlsx", "transforms"]
    template = (out / "attribute-mapping.csv").read_text(encoding="utf-8-sig").splitlines()
    assert template[0] == "old_attribute;new_attribute;status"
    assert "id;employeeId;from the mapping file" in template

    from openpyxl import load_workbook

    book = load_workbook(out / "source-migration.xlsx")
    assert book.sheetnames == ["Dashboard", "Next steps", "Profiles", "Attribute mapping", "Mappings",
                               "Transforms", "Criteria", "Checks", "Chart data", "Read me"]


@pytest.fixture
def criteria_plan():
    from isc_reporter.criteria import Objects

    return analyse(PROFILES, TRANSFORMS, Source("HR"), NEW, {"id": "employeeId"},
                   objects=Objects(ROLES, SEGMENTS, SOURCES))


def test_criteria_using_impacted_identity_attributes_or_the_old_source(criteria_plan):
    rows = {(r["Object"], r["Criterion"]): r for r in criteria_plan.criteria}
    department = rows[("Finance", 'identity.department EQUALS "Finance"')]
    assert (department["Severity"], department["Type"], department["Enabled"]) == ("Warning", "Role", "yes")
    assert department["Impact"].startswith(
        "Identity attribute computed from the old source (IP_Contractors):")
    assert department["All criteria"] == ('(identity.department EQUALS "Finance" AND identity.location '
                                          'EQUALS "Paris") OR HR.status EQUALS "A"')
    status = rows[("Finance", 'HR.status EQUALS "A"')]
    assert status["Uses"] == "HR account attribute 'status'"
    assert status["Action"] == "Rewrite on Workday: Workday.status"           # same name in the new schema
    old_group = rows[("Old HR group", 'entitlement HR.groups EQUALS "Staff"')]
    assert (old_group["Severity"], old_group["Enabled"]) == ("Info", "no")    # a disabled role
    leavers = rows[("Managers / Leavers", 'identity.cloudLifecycleState EQUALS "inactive"')]
    assert leavers["Type"] == "Role dimension"                                 # LCS comes from HR status
    segment = rows[("Finance visibility", 'identity.department EQUALS "Finance"')]
    assert (segment["Severity"], segment["All criteria"]) == (
        "Info", 'identity.department EQUALS "Finance" AND identity.location EQUALS "Paris"')
    ad = rows[("AD", "uid = sAMAccountName")]
    assert (ad["Severity"], ad["Type"], ad["Uses"]) == (
        "Warning", "Account correlation", "identity attribute 'uid'")
    assert rows[("HR", "uid = id")]["Action"].startswith("Recreate it on Workday")
    assert rows[("HR", "HR.managerId = identity.uid")]["Action"] == (
        "Recreate on Workday once 'managerId' has a counterpart in attribute-mapping.csv.")
    # attributes that are not impacted, identity lists and the new source itself are not reported
    assert {r["Object"] for r in criteria_plan.criteria} == {
        "Finance", "Old HR group", "Managers / Leavers", "Finance visibility", "AD", "HR"}
    assert not any("location" in r["Criterion"] or "email" in r["Criterion"] for r in criteria_plan.criteria)
    assert criteria_plan.criteria[0]["Severity"] == "Warning"


def test_old_source_attributes_of_criteria_join_the_attribute_mapping(criteria_plan, tmp_path):
    assert "role criteria: Finance" in criteria_plan.attributes["status"]["used_by"]
    assert criteria_plan.attributes["managerId"]["used_by"] == ["manager correlation"]
    checks = {(c[1], c[2]): c[0] for c in criteria_plan.checks}
    assert checks[("managerId", "New attribute to define")] == "Warning"
    assert checks[("Finance", "Role criteria to review")] == "Warning"
    assert checks[("Finance visibility", "Segment criteria to review")] == "Info"
    assert checks[("AD", "Account correlation criteria to review")] == "Warning"
    outside = next(c[3] for c in criteria_plan.checks if c[2] == "Outside this analysis")
    assert "role and segment criteria" not in outside

    from openpyxl import load_workbook

    from isc_reporter.source_migration import build_dashboard, write_outputs

    tiles = {label: value for label, value, _ in build_dashboard(criteria_plan).kpis}
    assert tiles["Criteria to review"] == "8  (3 role(s) or dimension(s))"
    write_outputs(criteria_plan, str(tmp_path), pdf=True)
    sheet = load_workbook(tmp_path / "source-migration.xlsx")["Criteria"]
    assert [c.value for c in sheet[1]] == ["Severity", "Type", "Object", "Enabled", "Criterion", "Uses",
                                           "Impact", "Action", "All criteria"]
    assert sheet.max_row == 9
    assert "managerId;;TO DEFINE" in (tmp_path / "attribute-mapping.csv").read_text(encoding="utf-8-sig")


def test_attributes_computed_from_impacted_ones_are_impacted_too():
    from isc_reporter.criteria import impacted_attributes

    profiles = [profile("IP_HR", "HR", mapping("uid", account("id")),
                        mapping("login", {"type": "identityAttribute", "attributes": {"name": "uid"}}),
                        mapping("alias", {"type": "static", "attributes": {"value": "$login-x"}}),
                        mapping("email", account("mail", "AD")))]
    assert impacted_attributes(profiles, {}, {"uid": ["IP_HR"]}) == {
        "uid": "computed from the old source (IP_HR)", "login": "computed from 'uid'",
        "alias": "computed from 'login'"}


def test_without_criteria_nothing_is_claimed(plan):
    from isc_reporter.criteria import Objects
    from isc_reporter.source_migration import build_dashboard

    assert not plan.criteria_analysed and plan.criteria == []
    assert "role and segment criteria" in next(c[3] for c in plan.checks if c[2] == "Outside this analysis")
    tiles = {label: value for label, value, _ in build_dashboard(plan).kpis}
    assert tiles["Criteria to review"] == "not analysed"
    # nothing could be read: the reasons are reported, and the analysis is not presented as done
    unread = analyse(PROFILES, TRANSFORMS, Source("HR"), NEW,
                     objects=Objects(notes=["Roles not read (HTTP 403: Forbidden)."]))
    assert not unread.criteria_analysed
    note = ("Info", "", "Criteria not fully analysed", "Roles not read (HTTP 403: Forbidden).")
    assert note in unread.checks


def test_criteria_objects_from_a_configuration_export(tmp_path):
    from isc_reporter.criteria import load_files

    managers = {k: v for k, v in ROLES[4].items() if k != "dimensions"}
    (tmp_path / "export.json").write_text(json.dumps({"version": 1, "objects": [
        {"self": {"type": "ROLE"}, "object": managers},
        {"self": {"type": "ROLE"}, "object": ROLES[0]},
        {"self": {"type": "DIMENSION"}, "object": LEAVERS},
        {"self": {"type": "SEGMENT"}, "object": SEGMENTS[0]},
        {"self": {"type": "SOURCE"}, "object": {**SOURCES[1], "connector": "delimited-file"}},
        {"self": {"type": "IDENTITY_PROFILE"}, "object": PROFILES[0]}]}), encoding="utf-8")
    objects = load_files([str(tmp_path)])
    assert [r["name"] for r in objects.roles] == ["Managers", "Finance"]
    assert objects.roles[0]["dimensions"] == [LEAVERS]
    assert [s["name"] for s in objects.segments] == ["Finance visibility"]
    assert [s["name"] for s in objects.sources] == ["HR"]
    assert objects.notes == []
    (tmp_path / "profiles.json").write_text(json.dumps(PROFILES), encoding="utf-8")
    only_profiles = load_files([str(tmp_path / "profiles.json")])
    assert only_profiles.roles is None and len(only_profiles.notes) == 3


def test_unknown_transforms_are_flagged():
    plan = analyse(PROFILES, {}, Source("HR"), NEW)
    flagged = {c[1] for c in plan.checks if c[2] == "Transform not analysed"}
    assert flagged == {"Format date", "LCS"}


def test_model_json_handles_dates_and_enums():
    from sailpoint.identity_profiles.models import IdentityProfile
    from sailpoint.sources.models.model_schema import ModelSchema

    from isc_reporter.api import model_json

    schema = ModelSchema.from_dict({"id": "s1", "name": "account", "nativeObjectType": "User",
                                    "created": "2024-03-17T09:31:59.493Z",
                                    "attributes": [{"name": "employeeId", "type": "STRING"}]})
    data = model_json(schema)
    assert data["created"].startswith("2024-03-17T09:31:59")
    assert (data["attributes"][0]["name"], data["attributes"][0]["type"]) == ("employeeId", "STRING")
    # The traceback case: a model with datetimes (some read-only, left out by the SDK).
    profile = IdentityProfile.from_dict({**PROFILES[0], "created": "2024-03-17T09:31:59.493Z",
                                         "modified": "2026-10-11T13:11:58.273Z"})
    assert model_json(profile)["name"] == "IP_HR"
    # str enums (str() gives "RoleCriteriaKeyType.IDENTITY") become their plain value
    from sailpoint.roles.models import Role

    role = Role.from_dict({**ROLES[1], "owner": OWNER})
    key = model_json(role)["membership"]["criteria"]["key"]
    assert type(key["type"]) is str and f"{key['type']}" == "IDENTITY"


def test_cli_against_tenant(tmp_path, monkeypatch):
    from sailpoint.identity_profiles.models import IdentityProfile
    from sailpoint.sources.models.model_schema import ModelSchema
    from sailpoint.transforms.models import Transform

    from conftest import FakeApi, FakeIsc

    dated = {"created": "2024-03-17T09:31:59.493Z", "modified": "2026-10-11T13:11:58.273Z"}
    profiles = [IdentityProfile.from_dict({**p, **dated}) for p in PROFILES]
    transforms = [Transform.from_dict(t) for t in TRANSFORMS.values()]
    sources = {"HR": "h" * 32, "Workday": "w" * 32}
    schemas = {"h" * 32: ["id", "endDate", "manager", "status", "department", "x"],
               "w" * 32: ["employeeId", "endDate", "manager", "status"]}

    from sailpoint.dimensions.models import Dimension
    from sailpoint.roles.models import Role
    from sailpoint.segments.models import Segment
    from sailpoint.sources.models import CorrelationConfig
    from sailpoint.sources.models import Source as SourceModel

    manager = {"accountAttributeName": "managerId", "identityAttributeName": "uid"}
    models = {name: SourceModel.from_dict({
        "id": source_id, "name": name, "owner": OWNER, "connector": "delimited-file",
        **({"managerCorrelationMapping": manager} if name == "HR" else {})})
        for name, source_id in sources.items()}
    correlation = {"h" * 32: CorrelationConfig.from_dict({"attributeAssignments": [assignment("uid", "id")]})}
    roles = [Role.from_dict({**finance("h" * 32), "owner": OWNER}),
             Role.from_dict({**{k: v for k, v in ROLES[4].items() if k != "dimensions"}, "owner": OWNER})]

    def list_sources_v1(filters=None, limit=250, offset=0):
        if filters:
            name = filters.split('"')[1]
            return [models[name]] if name in models else []
        return list(models.values()) if not offset else []

    def get_source_schemas_v1(source_id, include_types):
        return [ModelSchema.from_dict({"name": "account", "nativeObjectType": "User", **dated,
                                       "attributes": [{"name": n, "type": "STRING"}
                                                      for n in schemas[source_id]]})]

    def first_page(items):
        return lambda limit, offset: items if not offset else []

    isc = FakeIsc(
        identity_profiles=FakeApi(list_identity_profiles_v1=first_page(profiles)),
        transforms=FakeApi(list_transforms_v1=first_page(transforms)),
        identity_attributes=FakeApi(list_identity_attributes_v1=lambda **kw: []),
        sources=FakeApi(list_sources_v1=list_sources_v1, get_source_schemas_v1=get_source_schemas_v1,
                        get_correlation_config_v1=lambda id: correlation.get(id)),
        roles=FakeApi(list_roles_v1=first_page(roles)),
        dimensions=FakeApi(list_dimensions_v1=lambda role_id, limit, offset:
                           [Dimension.from_dict({**LEAVERS, "owner": OWNER})] if not offset else []),
        segments=FakeApi(list_segments_v1=first_page([Segment.from_dict(SEGMENTS[0])])))
    monkeypatch.setattr("isc_reporter.client.connect", lambda config=None: isc)
    out = tmp_path / "out"
    assert cli.main(["migration", "--old-source", "HR", "--new-source", "Workday", "--no-pdf",
                     "-o", str(out), "--log-dir", ""]) == 0
    template = (out / "attribute-mapping.csv").read_text(encoding="utf-8-sig")
    assert "endDate;endDate;same name in the new schema" in template
    assert "id;;TO DEFINE" in template
    assert isc.sources.called("get_source_schemas_v1")[0]["include_types"] == "account"
    assert not (out / "source-migration.pdf").exists()

    from openpyxl import load_workbook

    sheet = load_workbook(out / "source-migration.xlsx")["Criteria"]
    found = {(row[1], row[2], row[4]) for row in sheet.iter_rows(min_row=2, values_only=True)}
    assert ("Role", "Finance", 'HR.status EQUALS "A"') in found              # source id shown as its name
    assert ("Role dimension", "Managers / Leavers", 'identity.cloudLifecycleState EQUALS "inactive"') in found
    assert ("Segment", "Finance visibility", 'identity.department EQUALS "Finance"') in found
    assert ("Manager correlation", "HR", "HR.managerId = identity.uid") in found
    assert [kw["limit"] for kw in isc.roles.called("list_roles_v1")] == [50]   # the API allows 50 at most
    assert isc.dimensions.called("list_dimensions_v1")[0]["role_id"] == "managers"
