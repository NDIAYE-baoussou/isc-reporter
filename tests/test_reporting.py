"""Search reports: consolidation, summaries, Excel/PDF/CSV output and the command line."""

import json

import pytest

from conftest import FakeApi, FakeIsc
from isc_reporter import cli, report_writers, reporting
from isc_reporter.errors import UsageError
from isc_reporter.search import search_all


def access(kind, name, source=None, **extra):
    item = {"type": kind, "id": name.lower(), "name": name, "displayName": name, **extra}
    if source:
        item["source"] = {"name": source}
    return item


IDENTITIES = [
    {"id": "1", "name": "jdoe", "displayName": "John Doe", "email": "jdoe@example.com",
     "attributes": {"uid": "jdoe", "cloudLifecycleState": "active", "department": "Finance"},
     "identityProfile": {"name": "HR"}, "manager": {"id": "9", "displayName": "Ann Boss"},
     "created": "2026-01-10T08:00:00Z", "accessCount": 4, "accountCount": 2,
     "access": [access("ROLE", "Finance Analyst"), access("ACCESS_PROFILE", "Finance - Read"),
                access("ENTITLEMENT", "APP-Finance-Read", "AD"), access("ENTITLEMENT", "Users", "AD")],
     "accounts": [{"name": "jdoe", "source": {"name": "AD"}},
                  {"name": "jdoe@x", "source": {"name": "Mail"}, "disabled": True}]},
    {"id": "2", "name": "asmith", "displayName": "Alice Smith", "email": "=cmd()",
     "attributes": {"uid": "asmith", "cloudLifecycleState": "inactive", "department": "IT"},
     "identityProfile": {"name": "HR"}, "inactive": True, "created": "2026-01-11T08:00:00Z",
     "accessCount": 30, "accountCount": 0, "access": [access("ROLE", "Finance Analyst")], "accounts": []},
    {"id": "3", "name": "bob", "displayName": "Bob", "attributes": {"uid": "bob"},
     "created": "2026-01-11T09:00:00Z", "accessCount": 0},
]


@pytest.fixture
def report():
    return reporting.build(reporting.report_type("identities"), IDENTITIES, query="*", expand="access",
                           origin="https://tenant")


def test_one_line_per_identity_with_consolidated_access(report):
    rows = report.rows()
    assert len(rows) == 3
    john = dict(zip(report.headers, rows[0], strict=True))
    assert (john["Identity"], john["Username"], john["Manager"]) == ("John Doe", "jdoe", "Ann Boss")
    assert (john["Roles"], john["Access profiles"], john["Entitlements"], john["Accounts"]) == (1, 1, 2, 2)
    assert john["Entitlement list"] == "APP-Finance-Read (AD) | Users (AD)"
    assert john["Account list"] == "AD: jdoe | Mail: jdoe@x (disabled)"
    bob = dict(zip(report.headers, rows[2], strict=True))
    assert (bob["Roles"], bob["Role list"], bob["Manager"]) == (0, "", "")


def test_detail_rows_one_per_access_item(report):
    rows = report.detail_rows()
    assert report.detail_headers()[:4] == ["Identity", "Username", "Email", "Access type"]
    assert len(rows) == 4 + 1 + 1  # Bob has no access but keeps a line
    assert rows[2][3:6] == ["ENTITLEMENT", "APP-Finance-Read (AD)", "AD"]
    assert rows[-1][:2] == ["Bob", "bob"] and rows[-1][3] == ""


def test_breakdowns_and_figures(report):
    by_title = {b.title: report.breakdown(b) for b in report.breakdowns}
    assert by_title["Lifecycle state"] == [("active", 1), ("inactive", 1), ("(empty)", 1)]
    assert by_title["Access items per identity"] == [("0", 1), ("1-5", 1), ("6-20", 0), ("21-50", 1),
                                                      ("51-100", 0), ("> 100", 0)]
    assert by_title["Most held roles"] == [("Finance Analyst", 2)]
    figures = dict(report.figures())
    assert figures["Results"] == 3
    assert figures["Inactive identities"] == "1 (33%)"
    assert figures["Identities without a manager"] == "2 (67%)"


def test_top_categories_and_dates():
    docs = [{"type": f"T{i % 20}", "created": f"2026-02-{1 + i % 3:02d}T10:00:00Z"} for i in range(40)]
    rt = reporting.report_type("events")
    report = reporting.build(rt, docs)
    by_title = {b.title: report.breakdown(b) for b in report.breakdowns}
    types = by_title["Type"]
    assert len(types) == reporting.TOP + 1 and types[-1] == ("(other)", 10)
    assert by_title["Events per day"] == [("2026-02-01", 14), ("2026-02-02", 13), ("2026-02-03", 13)]


def test_custom_columns_group_by_and_generic_index():
    rt = reporting.report_type("identities")
    report = reporting.build(rt, IDENTITIES, columns=["displayName", "attributes.department", "access.name"],
                             group_by=["attributes.department"])
    assert report.headers == ["Display name", "Department", "Name"]
    assert report.rows()[0] == ["John Doe", "Finance",
                                "Finance Analyst | Finance - Read | APP-Finance-Read | Users"]
    assert report.breakdown(report.breakdowns[0])[0] == ("Finance", 1)
    generic = reporting.build(reporting.report_type("*"), [{"id": "x", "name": "n", "meta": {"k": 1}}])
    assert generic.headers[:2] == ["Name", "Id"] and "K" in generic.headers
    with pytest.raises(UsageError, match="cannot expand"):
        reporting.build(rt, IDENTITIES, expand="nope")
    with pytest.raises(UsageError):
        reporting.report_type("unknown")


def test_definition_file(tmp_path):
    path = tmp_path / "r.json"
    path.write_text(json.dumps({"type": "roles", "query": "enabled:false", "group_by": ["owner.name"]}))
    assert reporting.load_definition(str(path))["query"] == "enabled:false"
    path.write_text(json.dumps({"type": "roles", "colums": []}))
    with pytest.raises(UsageError, match="colums"):
        reporting.load_definition(str(path))


def test_excel_workbook(report, tmp_path):
    from openpyxl import load_workbook

    path = report_writers.write_xlsx(report, str(tmp_path / "r.xlsx"))
    book = load_workbook(path)
    assert book.sheetnames == ["Summary", "Data", "Detail - Access", "Query"]
    data = book["Data"]
    assert "Results" in data.tables and data.max_row == 4
    assert data["C3"].value == "=cmd()" and data["C3"].data_type == "s"   # never a formula
    assert data["J2"].value == 1                                         # counts stay numbers
    summary = book["Summary"]
    assert len(summary._charts) == len(report.breakdowns)
    assert book["Detail - Access"].max_row == 7
    assert book["Query"]["B3"].value == "*"


def test_pdf_and_csv(report, tmp_path):
    pdf = report_writers.write_pdf(report, str(tmp_path / "r.pdf"), max_rows=2)
    assert open(pdf, "rb").read(4) == b"%PDF"
    paths = report_writers.write_csv(report, str(tmp_path / "r.csv"))
    assert [p.split("r")[-1] for p in paths] == [".csv", "-access.csv"]
    assert open(paths[0], encoding="utf-8-sig").readline().startswith("Identity;Username;Email")


def test_search_all_limit_and_id_sort():
    bodies = []

    def search_post_v1(search, limit):
        bodies.append((search.to_dict(), limit))
        return [{"id": str(i)} for i in range(limit)]

    isc = FakeIsc(search=FakeApi(search_post_v1=search_post_v1))
    docs = search_all(isc, "roles", "*", sort_field=None, max_results=300)
    assert len(docs) == 300
    assert bodies[0][0]["sort"] == ["id"] and [b[1] for b in bodies] == [250, 50]
    assert bodies[1][0]["searchAfter"] == ["249"]


def test_cli_report_search(tmp_path, monkeypatch):
    isc = FakeIsc(search=FakeApi(search_post_v1=lambda search, limit: IDENTITIES[:limit]))
    monkeypatch.setattr("isc_reporter.client.connect", lambda config=None: isc)
    out = tmp_path / "report"
    assert cli.main(["search", "identities", "-q", "attributes.department:Finance",
                     "--expand", "accounts", "--format", "xlsx", "pdf", "csv", "--limit", "2",
                     "-o", str(out), "--log-dir", ""]) == 0
    assert {p.name for p in tmp_path.iterdir()} == {"report.xlsx", "report.pdf", "report.csv",
                                                    "report-accounts.csv"}
    query = isc.search.called("search_post_v1")[0]["search"].to_dict()
    assert query["indices"] == ["identities"] and query["query"]["query"] == "attributes.department:Finance"
    assert cli.main(["search", "--log-dir", ""]) == 2


def test_example_definitions_are_valid():
    import glob

    paths = glob.glob("examples/reports/*.json")
    assert paths
    for path in paths:
        definition = reporting.load_definition(path)
        rt = reporting.report_type(definition["type"])
        reporting.build(rt, IDENTITIES if rt.name == "identities" else [], expand=definition.get("expand"),
                        group_by=definition.get("group_by"))


def test_excel_package_is_what_excel_itself_writes(report, tmp_path):
    """Shared strings and cached chart categories: rendered by every viewer, not only Excel."""
    import re
    import zipfile

    path = report_writers.write_xlsx(report, str(tmp_path / "r.xlsx"))
    package = zipfile.ZipFile(path)
    assert "xl/sharedStrings.xml" in package.namelist()
    assert "inlineStr" not in package.read("xl/worksheets/sheet2.xml").decode()
    chart = package.read("xl/charts/chart1.xml").decode()
    assert re.search(r"<c:cat><c:strRef>.*<c:strCache>", chart)
    assert "<f>" not in package.read("xl/worksheets/sheet2.xml").decode()   # '=cmd()' stays text
