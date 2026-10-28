"""Forecast of the updates caused by enabling attribute synchronization."""

import pytest
from openpyxl import load_workbook

from conftest import FakeApi, FakeIsc
from isc_reporter import cli, sync_impact
from isc_reporter.errors import UsageError


def direct(name):
    return {"type": "identityAttribute", "attributes": {"name": name}}


POLICIES = [{"name": "AD Create", "usageType": "CREATE", "fields": [
    {"name": "mail", "transform": direct("email")},
    {"name": "department", "transform": direct("department")},
    {"name": "title", "transform": direct("title")},
    {"name": "cn", "transform": {"type": "rule", "attributes": {"name": "Build CN"}}},
    {"name": "info"}]},
    {"name": "AD Update", "usageType": "UPDATE", "fields": [{"name": "sn", "transform": direct("lastname")}]}]


def account(n, identity, **attributes):
    return {"id": f"a{n}", "name": f"acc{n}", "identityId": identity, "attributes": attributes,
            "disabled": False}


ACCOUNTS = [
    account(1, "i1", mail="a@x.com", department="Sales", title="Boss"),         # unchanged
    account(2, "i2", mail="OLD@x.com", department="Sales"),                      # change, add title
    account(3, "i3", mail="c@x.com", department="Ops", title="Dev"),             # title cleared
    account(4, "i4", mail="D@X.COM", department="Sales", title="Dev"),          # case only
    {"id": "a5", "name": "acc5", "identityId": None, "uncorrelated": True, "attributes": {}},
    {**account(6, "i9", mail="z"), "disabled": True},                           # identity unreadable
]


def identity(n, name, **attributes):
    return {"id": f"i{n}", "displayName": name, "attributes": attributes}


IDENTITIES = {
    "i1": identity(1, "One", email="a@x.com", department="Sales", title="Boss"),
    "i2": identity(2, "Two", email="b@x.com", department="Sales", title="CEO"),
    "i3": identity(3, "Three", email="c@x.com", department="Ops"),
    "i4": identity(4, "Four", email="d@x.com", department="Sales", title="Dev"),
}


def test_policy_mappings_keep_only_direct_copies_of_the_wanted_usage():
    mappings, ignored = sync_impact.policy_mappings(POLICIES)
    assert [(m.identity_attribute, m.account_attribute) for m in mappings] == \
        [("email", "mail"), ("department", "department"), ("title", "title")]
    assert [(i["Field"], i["Transform"]) for i in ignored] == [("cn", "rule"), ("info", "none")]
    both, _ = sync_impact.policy_mappings(POLICIES, ("UPDATE", "CREATE"))
    assert len(both) == 4


def test_maps_from_the_command_line_and_a_file(tmp_path):
    parsed = sync_impact.parse_maps(["a=b", " c = d "])
    assert [(m.identity_attribute, m.account_attribute) for m in parsed] == [("a", "b"), ("c", "d")]
    with pytest.raises(UsageError):
        sync_impact.parse_maps(["nope"])
    path = tmp_path / "map.csv"
    path.write_text("identity_attribute;account_attribute\nlocation;physicalDeliveryOfficeName\n",
                    encoding="utf-8")
    assert sync_impact.read_map_file(str(path))[0].account_attribute == "physicalDeliveryOfficeName"
    policy, _ = sync_impact.policy_mappings(POLICIES)
    merged = sync_impact.merge(policy, sync_impact.parse_maps(["EMAIL=MAIL", "location=l"]))
    assert len(merged) == 4 and merged[0].origin.startswith("CREATE")      # same pair: the policy keeps it


def test_classify():
    values = sync_impact.values_of
    assert sync_impact.classify(values("a"), values("a")) is None
    assert sync_impact.classify(values(None), values("")) is None
    assert sync_impact.classify(values(["b", "a"]), values(["a", "b"])) is None
    assert sync_impact.classify(values(""), values("x")) == sync_impact.ADD
    assert sync_impact.classify(values("x"), values(None)) == sync_impact.CLEAR
    assert sync_impact.classify(values("x"), values("X")) == sync_impact.CASE
    assert sync_impact.classify(values("x"), values("y")) == sync_impact.CHANGE
    assert sync_impact.values_of(True) == ("true",)


def test_analyse_counts_updates_per_account_and_attribute():
    mappings, ignored = sync_impact.policy_mappings(POLICIES)
    result = sync_impact.analyse("AD", ACCOUNTS, IDENTITIES, mappings, ignored, threshold=50)
    assert (result.accounts, result.analysed, result.uncorrelated, result.identity_missing) == (6, 4, 1, 1)
    kinds = {m.account_attribute: dict(s.kinds) for m, s in result.stats.items()}
    assert kinds == {"mail": {"Change value": 1, "Case only": 1}, "department": {},
                     "title": {"Add value": 1, "Clear value": 1}}
    assert result.outcomes == {"a2": "Change value", "a3": "Clear value", "a4": "Case only"}
    assert result.impacted == 3 and result.verdict == "Bulk update expected"        # a value is cleared
    found = {(c[1], c[2]) for c in sync_impact.checks(result)}
    assert ("email -> mail", "Case only") in found
    assert ("title -> title", "Values would be cleared") in found
    assert ("", "Uncorrelated accounts") in found
    assert ("", "Identity not found") in found
    assert ("", "Fields not analysed") in found
    cleared = next(c for c in result.changes if c.kind == "Clear value")
    assert (cleared.account, cleared.identity, cleared.current, cleared.new) == ("acc3", "Three", "Dev", "")
    assert cleared.mapping.account_attribute == "title"


def test_no_difference_means_no_update():
    mappings = sync_impact.parse_maps(["email=mail"])
    result = sync_impact.analyse("AD", ACCOUNTS[:1], IDENTITIES, mappings)
    assert result.verdict == "No update expected" and result.changes == []


def test_absent_attributes_are_reported():
    result = sync_impact.analyse("AD", ACCOUNTS[:2], IDENTITIES, sync_impact.parse_maps(["ghost=nope"]))
    checks = {c[2] for c in sync_impact.checks(result)}
    assert {"Account attribute not found", "Identity attribute not found"} <= checks


def test_reports(tmp_path):
    mappings, ignored = sync_impact.policy_mappings(POLICIES)
    result = sync_impact.analyse("AD", ACCOUNTS, IDENTITIES, mappings, ignored)
    xlsx, pdf, csv = (str(tmp_path / n) for n in ("r.xlsx", "r.pdf", "r.csv"))
    sync_impact.write_xlsx(xlsx, result, "tenant")
    sync_impact.write_pdf(pdf, result, "tenant")
    sync_impact.write_csv(csv, result)
    book = load_workbook(xlsx)
    assert book.sheetnames == ["Dashboard", "Attributes", "Accounts", "Changes", "Not analysed", "Checks",
                               "Chart data", "Read me"]
    rows = list(book["Attributes"].iter_rows(min_row=2, values_only=True))
    assert [r[:2] for r in rows] == [("email", "mail"), ("department", "department"), ("title", "title")]
    assert rows[2][-1] == "High"                                    # a value would be cleared
    assert open(pdf, "rb").read(4) == b"%PDF"
    lines = open(csv, encoding="utf-8-sig").read().splitlines()
    assert lines[0].startswith("Account;Identity") and len(lines) == 1 + len(result.changes)


def fake_isc(accounts=ACCOUNTS, policies=POLICIES):
    from sailpoint.sources.models import ProvisioningPolicyDto

    source = type("S", (), {"id": "s1", "name": "AD"})()
    models = [ProvisioningPolicyDto.from_dict(p) for p in policies]
    return FakeIsc(
        sources=FakeApi(list_sources_v1=lambda **kw: [source],
                        list_provisioning_policies_v1=lambda **kw: models if kw.get("offset") == 0 else []),
        accounts=FakeApi(list_accounts_v1=lambda **kw: accounts if kw.get("offset") == 0 else []),
        search=FakeApi(search_post_v1=lambda **kw: list(IDENTITIES.values())))


def test_cli(tmp_path, monkeypatch):
    isc = fake_isc()
    monkeypatch.setattr("isc_reporter.client.connect", lambda config=None: isc)
    out = tmp_path / "impact"
    args = ["analyzebeforesync", "AD", "--format", "xlsx", "csv", "-o", str(out), "--log-dir", ""]
    assert cli.main(args) == 0
    assert (tmp_path / "impact.xlsx").exists() and (tmp_path / "impact.csv").exists()
    assert isc.accounts.called("list_accounts_v1")[0]["filters"] == 'sourceId eq "s1"'
    query = isc.search.called("search_post_v1")[0]["search"].query.query
    assert query.startswith("id:(") and '"i1"' in query and "None" not in query


def test_cli_without_any_direct_mapping_asks_for_a_map(tmp_path, monkeypatch):
    monkeypatch.setattr("isc_reporter.client.connect", lambda config=None: fake_isc(policies=[]))
    base = ["analyze-before-sync", "AD", "--format", "csv", "-o", str(tmp_path / "x"), "--log-dir", ""]
    assert cli.main(base) == 2
    assert cli.main([*base, "--map", "email=mail"]) == 0
    assert cli.main([*base, "--map", "email=mail", "--only-map"]) == 0
