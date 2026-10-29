"""Leavers who still have access."""

from datetime import date

from openpyxl import load_workbook

from conftest import FakeApi, FakeIsc
from isc_reporter import cli, leavers

AD, SAP, CRM, HR = "a" * 32, "s" * 32, "c" * 32, "h" * 32


def action(kind, sources=None, all_sources=False, exclude=None):
    return {"action": kind, "sourceIds": sources or [], "allSources": all_sources,
            "excludeSourceIds": exclude or []}


def state(name, identity_state, *actions):
    return {"name": name, "technicalName": name, "identityState": identity_state, "enabled": True,
            "accountActions": list(actions)}


LIFECYCLE = {"p1": {
    "active": state("active", "ACTIVE"),
    "inactive": state("inactive", "INACTIVE_SHORT_TERM", action("DISABLE", all_sources=True, exclude=[CRM])),
    "out": state("out", "INACTIVE_LONG_TERM", action("DELETE", [AD])),
}}


def account(source_id, source_name, disabled=False, privileged=False):
    name = f"{source_name.lower()}-account"
    return {"id": name, "name": name, "accountId": name, "source": {"id": source_id, "name": source_name},
            "disabled": disabled, "privileged": privileged}


def identity(name, lcs, end, accounts, profile="p1"):
    return {"id": name, "name": name, "displayName": name.title(),
            "identityProfile": {"id": profile, "name": "Employees"},
            "attributes": {"cloudLifecycleState": lcs, "endDate": end}, "accounts": accounts}


IDENTITIES = [
    identity("alice", "inactive", "2026-10-20", [account(AD, "AD"), account(CRM, "CRM"), account(HR, "HR")]),
    identity("bob", "inactive", "2026-10-28", [account(AD, "AD", disabled=True)]),
    identity("carol", "out", "2026-01-15T00:00:00Z",
             [account(AD, "AD", privileged=True), account(SAP, "SAP")]),
    identity("dan", "gardenLeave", "15/10/2026", [account(SAP, "SAP")], profile="p2"),
]


def data(**changes):
    values = {"states": ["inactive", "out", "gardenLeave"], "identities": IDENTITIES, "lifecycle": LIFECYCLE,
              "profiles": {"p1": "Employees"}, "skipped_sources": {HR: "authoritative source"},
              "date_attribute": "endDate", "grace": 5, "today": date(2026, 10, 29)}
    return leavers.LeaverData(**{**values, **changes})


def test_default_states_are_the_inactive_ones():
    assert leavers.default_states(LIFECYCLE) == ["inactive", "out"]
    assert leavers.default_states({"p1": {"active": state("active", "ACTIVE")}}) == []


def test_expected_action_follows_the_state_configuration():
    inactive, out = LIFECYCLE["p1"]["inactive"], LIFECYCLE["p1"]["out"]
    assert leavers.expected_action(inactive, AD) == "DISABLE"            # all sources...
    assert leavers.expected_action(inactive, CRM) is None                # ...except CRM
    assert leavers.expected_action(out, AD) == "DELETE"
    assert leavers.expected_action(out, SAP) is None
    assert leavers.cause(None, AD) == leavers.UNKNOWN
    assert leavers.cause(state("x", "ACTIVE", action("ENABLE", [AD])), AD) == leavers.KEPT


def test_departure_dates():
    assert leavers.left_on("2026-10-20T08:00:00Z") == date(2026, 10, 20)
    assert leavers.left_on("15/10/2026") == leavers.left_on("20261015") == date(2026, 10, 15)
    assert leavers.left_on("soon") is None and leavers.left_on("") is None


def test_enabled_accounts_and_their_cause():
    result = leavers.analyse(data())
    assert (result.total, len(result.leavers), result.without_access) == (4, 3, 1)       # bob is clean
    rows = [(r["Identity"], r["Source"], r["Cause"], r["Days since"], r["Beyond grace"], r["Privileged"])
            for r in result.accounts]
    assert rows == [
        ("carol", "AD", leavers.DELETE, 287, "yes", "yes"),                 # privileged first
        ("alice", "AD", leavers.DISABLE, 9, "yes", "no"),
        ("carol", "SAP", leavers.NOT_COVERED, 287, "yes", "no"),          # longest gone first
        ("alice", "CRM", leavers.NOT_COVERED, 9, "yes", "no"),
        ("dan", "SAP", leavers.UNKNOWN, 14, "yes", "no"),                   # profile p2 not read
    ]
    assert "HR" not in {r["Source"] for r in result.accounts}               # authoritative: left out
    assert result.leavers[0]["Identity"] == "carol" and result.leavers[0]["Privileged"] == 1
    within = leavers.analyse(data(grace=10))
    assert {r["Identity"]: r["Beyond grace"] for r in within.accounts}["alice"] == "no"


def test_without_date_attribute_and_unreadable_dates():
    result = leavers.analyse(data(date_attribute=""))
    assert {r["Days since"] for r in result.accounts} == {None}
    assert {r["Beyond grace"] for r in result.accounts} == {""}
    odd = [identity("eve", "inactive", "next week", [account(AD, "AD")])]
    assert leavers.analyse(data(identities=odd)).unparsed_dates == 1


def test_reports(tmp_path):
    result = leavers.analyse(data())
    xlsx, pdf, csv = (str(tmp_path / f"l.{ext}") for ext in ("xlsx", "pdf", "csv"))
    leavers.write_xlsx(xlsx, result)
    leavers.write_pdf(pdf, result)
    leavers.write_csv(csv, result)
    book = load_workbook(xlsx)
    assert book.sheetnames == ["Dashboard", "Leavers with access", "Enabled accounts", "By source",
                               "Lifecycle states", "Settings", "Chart data", "Read me"]
    states = {(r[0], r[1]): r for r in book["Lifecycle states"].iter_rows(min_row=2, values_only=True)}
    assert states[("Employees", "inactive")][3] == "DISABLE on all sources except 1"
    assert states[("Employees", "inactive")][4] == 2
    by_source = {r[0]: r for r in book["By source"].iter_rows(min_row=2, values_only=True)}
    assert by_source["AD"][1:3] == (2, 1)                                   # enabled, privileged
    assert open(pdf, "rb").read(4) == b"%PDF"
    lines = open(csv, encoding="utf-8-sig").read().splitlines()
    assert lines[0].startswith("Identity;Display name;Lifecycle state") and len(lines) == 6


def fake_isc(lifecycle=LIFECYCLE):
    from sailpoint.lifecycle_states.models import LifecycleState

    profiles = [type("P", (), {"id": pid, "name": "Employees"})() for pid in lifecycle]
    sources = [type("S", (), {"id": sid, "name": name, "authoritative": name == "HR"})()
               for sid, name in ((AD, "AD"), (SAP, "SAP"), (CRM, "CRM"), (HR, "HR"))]

    def states(identity_profile_id, limit, offset):
        found = lifecycle.get(identity_profile_id, {}).values()
        return [LifecycleState.from_dict(s) for s in found] if not offset else []

    return FakeIsc(
        identity_profiles=FakeApi(list_identity_profiles_v1=lambda limit, offset:
                                  profiles if not offset else []),
        lifecycle_states=FakeApi(get_lifecycle_states_v1=states),
        sources=FakeApi(list_sources_v1=lambda limit, offset: sources if not offset else []),
        search=FakeApi(search_post_v1=lambda **kw: IDENTITIES))


def test_cli(tmp_path, monkeypatch):
    isc = fake_isc()
    monkeypatch.setattr("isc_reporter.client.connect", lambda config=None: isc)
    out = tmp_path / "leavers"
    assert cli.main(["leavers", "--date-attribute", "endDate", "--grace", "2", "--format", "xlsx", "csv",
                     "-o", str(out), "--log-dir", ""]) == 0
    query = isc.search.called("search_post_v1")[0]["search"].query.query
    assert query == 'attributes.cloudLifecycleState:("inactive" OR "out")'      # inactive states by default
    rows = list(load_workbook(str(out) + ".xlsx")["Enabled accounts"].iter_rows(min_row=2, values_only=True))
    assert {(r[0], r[5], r[8]) for r in rows} >= {("alice", "AD", leavers.DISABLE),
                                                   ("alice", "CRM", leavers.NOT_COVERED)}
    assert (tmp_path / "leavers.csv").exists()


def test_cli_needs_states_or_a_date_for_grace(tmp_path, monkeypatch):
    monkeypatch.setattr("isc_reporter.client.connect",
                        lambda config=None: fake_isc({"p1": {"active": state("active", "ACTIVE")}}))
    base = ["leavers", "--format", "csv", "-o", str(tmp_path / "x"), "--log-dir", ""]
    assert cli.main(base) == 2                                    # no inactive state configured
    assert cli.main([*base, "--grace", "3"]) == 2                 # grace without a departure date
    assert cli.main([*base, "--states", "inactive", "gardenLeave"]) == 0
