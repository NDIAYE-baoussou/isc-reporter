"""Identity attribute change history (old and new values)."""

import json
import zipfile
from datetime import datetime, timezone

import pytest

from conftest import FakeApi, FakeIsc
from isc_reporter import attribute_history as ah
from isc_reporter import cli
from isc_reporter.errors import UsageError

START = datetime(2026, 9, 27, tzinfo=timezone.utc)
END = datetime(2026, 10, 27, tzinfo=timezone.utc)
BOSS_ID = "b" * 32


def identity(iid, first, last, number):
    return {"id": iid, "name": first.lower(), "displayName": f"{first} {last}", "firstName": first,
            "lastName": last, "email": f"{first.lower()}@example.com",
            "attributes": {"employeeNumber": number}}


JOHN = identity("1" * 32, "John", "Doe", "E001")
ANN = identity("2" * 32, "Ann", "Lee", "E002")


def event(when, *changes, kind="AttributesChanged"):
    return {"eventType": kind, "dateTime": when,
            "attributeChanges": [{"name": n, "previousValue": o, "newValue": v} for n, o, v in changes]}


EVENTS = {
    JOHN["id"]: [
        event("2026-10-02T09:00:00Z", ("department", "Finance", "IT"), ("manager", "Ann Boss", BOSS_ID)),
        event("2026-10-20T09:00:00Z", ("department", "IT", "Security"), ("location", "Paris", "Lyon")),
        event("2026-08-01T09:00:00Z", ("department", "HR", "Finance")),          # before the period
        event("2026-10-05T09:00:00Z", kind="AccessItemAssociated"),             # other event type
    ],
    ANN["id"]: [event("2026-10-10T09:00:00Z", ("manager", {"id": "x", "name": "Old Boss"}, "New Boss"))],
}


@pytest.fixture
def data():
    d = ah.HistoryData(START, END, ["department", "manager"], identities=[JOHN, ANN])
    wanted = {"department", "manager"}
    for person in (JOHN, ANN):
        d.changes += ah.changes_from(person, EVENTS[person["id"]], START, END, wanted)
    return d


def test_changes_are_filtered_and_values_readable(data):
    rows = ah.change_rows(data)
    assert [(r["Identity"], r["Attribute"], r["Old value"], r["New value"]) for r in rows] == [
        ("John Doe", "Department", "Finance", "IT"), ("John Doe", "Manager", "Ann Boss", BOSS_ID),
        ("John Doe", "Department", "IT", "Security"), ("Ann Lee", "Manager", "Old Boss", "New Boss")]
    assert rows[0]["Employee number"] == "E001" and rows[0]["Date"] == "2026-10-02 09:00"


def test_one_line_per_identity_in_the_requested_layout(data):
    rows = {r["Identity"]: r for r in ah.identity_rows(data)}
    john = rows["John Doe"]
    assert list(john)[:8] == ["Employee number", "First name", "Last name", "Email", "Identity",
                              "Old Department", "New Department", "Old Manager"]
    # first old value and last new value of the period
    assert (john["Old Department"], john["New Department"]) == ("Finance", "Security")
    assert (john["Changes"], john["Last change"]) == (3, "2026-10-20")
    ann = rows["Ann Lee"]
    assert (ann["Old Department"], ann["New Department"], ann["New Manager"]) == ("", "", "New Boss")


def test_all_attributes_when_none_requested():
    d = ah.HistoryData(START, END, None, identities=[JOHN])
    d.changes = ah.changes_from(JOHN, EVENTS[JOHN["id"]], START, END, None)
    assert ah.changed_attributes(d) == ["department", "manager", "location"]
    assert "Old Location" in ah.identity_rows(d)[0]


def test_summaries_and_charts(data):
    assert ah.attribute_rows(data)[0] == {"Attribute": "Department", "Technical name": "department",
                                          "Changes": 2, "Identities": 1, "Days with changes": 2,
                                          "Note": ""}
    assert {(t["From"], t["To"]) for t in ah.transition_rows(data)} >= {("Finance", "IT"), ("IT", "Security")}
    tiles = {label: value for label, value, _ in ah.kpis(data)}
    assert tiles["Identities with changes"] == "2  (of 2 modified)"
    assert tiles["Most changed attribute"] in ("Department (2)", "Manager (2)")
    titles = [c.title for c in ah.charts(data)]
    assert titles[:2] == ["Changes per attribute", "Changes per day"]
    assert "Top Department changes" in titles and "Top Manager changes" in titles


def test_labels_and_period():
    assert ah.label("cloudLifecycleState") == "Lifecycle State"
    assert ah.label("costCenter") == "Cost Center"
    start, end = ah.period(30, now=END)
    assert (end - start).days == 30
    assert ah.period(since="2026-10-01", now=END)[0] == datetime(2026, 10, 1, tzinfo=timezone.utc)
    with pytest.raises(UsageError):
        ah.period(since="01/10/2026")


def test_outputs(data, tmp_path):
    from openpyxl import load_workbook

    xlsx = ah.write_xlsx(data, str(tmp_path / "r.xlsx"), "Changes")
    book = load_workbook(xlsx)
    assert book.sheetnames == ["Dashboard", "By identity", "Changes", "Attributes", "Transitions",
                               "Chart data", "Read me"]
    assert "xl/sharedStrings.xml" in zipfile.ZipFile(xlsx).namelist()
    assert open(ah.write_pdf(data, str(tmp_path / "r.pdf"), "Changes"), "rb").read(4) == b"%PDF"
    paths = ah.write_csv(data, str(tmp_path / "r"))
    header = open(paths[0], encoding="utf-8-sig").readline().strip()
    assert header.startswith("Employee number;First name;Last name;Email;Identity;Old Department")
    empty = ah.HistoryData(START, END, ["department"])
    ah.write_xlsx(empty, str(tmp_path / "e.xlsx"), "Empty")
    ah.write_pdf(empty, str(tmp_path / "e.pdf"), "Empty")


class Raw:
    """What a ``..._without_preload_content`` call returns."""

    def __init__(self, payload, status=200):
        self.status, self.headers = status, {}
        self.data = json.dumps(payload).encode()


def tenant(history_status=200):
    def search_post_v1(search, limit):
        query = search.query.query
        if query.startswith("id:"):                       # manager id resolution
            return [{"id": BOSS_ID, "displayName": "Bob Chief"}]
        return [JOHN, ANN] if not search.search_after else []

    def history(id, event_types, limit, offset, x_sail_point_experimental):
        assert event_types == ["AttributesChanged"] and x_sail_point_experimental == "true"
        return Raw(EVENTS[id] if not offset else [], history_status)

    return FakeIsc(search=FakeApi(search_post_v1=search_post_v1),
                   identity_history=FakeApi(get_historical_identity_events_v1_without_preload_content=history))


def test_fetch_reads_history_and_resolves_ids():
    isc = tenant()
    d = ah.fetch(isc, START, END, ["department", "manager"], query="attributes.department:*")
    query = isc.search.called("search_post_v1")[0]["search"].query.query
    assert query == "(attributes.department:*) AND modified:[2026-09-27T00:00:00.000Z TO now]"
    managers = [c["new"] for c in d.changes if c["attribute"] == "manager"]
    assert managers == ["Bob Chief", "New Boss"]


def test_missing_permission_is_explained():
    with pytest.raises(UsageError, match="identity-history:read"):
        ah.fetch(tenant(history_status=403), START, END)


def test_cli(tmp_path, monkeypatch):
    monkeypatch.setattr("isc_reporter.client.connect", lambda config=None: tenant())
    out = tmp_path / "report"
    assert cli.main(["changes", "--attributes", "department,manager", "--days", "3650",
                     "--format", "xlsx", "pdf", "csv", "-o", str(out), "--log-dir", ""]) == 0
    assert {p.name for p in tmp_path.iterdir()} == {"report.xlsx", "report.pdf", "report.csv",
                                                    "report-changes.csv"}


def test_daily_recomputed_attributes_are_flagged_and_can_be_excluded():
    days = [f"2026-10-{d:02d}T19:00:00Z" for d in range(1, 27)]
    events = [event(day, ("action", f"enabled on 2026-10-{i:02d}", f"enabled on 2026-10-{i + 1:02d}"))
              for i, day in enumerate(days)] + [event("2026-10-02T09:00:00Z", ("department", "HR", "IT"))]
    d = ah.HistoryData(START, END, None, identities=[JOHN])
    d.changes = ah.changes_from(JOHN, events, START, END, None)
    assert ah.daily_attributes(d) == {"action"}
    notes = {r["Technical name"]: r["Note"] for r in ah.attribute_rows(d)}
    assert notes["action"].startswith("Changes on most days") and notes["department"] == ""
    assert ("Recomputed every day (see Attributes)", "Action") in [(k, v) for k, v, _ in ah.kpis(d)]
    kept = ah.changes_from(JOHN, events, START, END, None, excluded={"action"})
    assert [c["attribute"] for c in kept] == ["department"]


def test_long_labels_keep_their_distinct_ending():
    label = ah._short("Account enabled by Identity Security Cloud (ISC) Automation on 2026-09-26 → "
                      "Account enabled by Identity Security Cloud (ISC) Automation on 2026-09-27")
    assert len(label) == 45 and label.endswith("2026-09-27") and "…" in label


def test_csv_keeps_its_columns_without_changes(tmp_path):
    empty = ah.HistoryData(START, END, ["department", "manager"])
    by_identity, changes = ah.write_csv(empty, str(tmp_path / "e"))
    assert open(by_identity, encoding="utf-8-sig").read().strip() == (
        "Employee number;First name;Last name;Email;Identity;Old Department;New Department;"
        "Old Manager;New Manager;Changes;Last change")
    assert open(changes, encoding="utf-8-sig").readline().startswith("Date;Employee number")
