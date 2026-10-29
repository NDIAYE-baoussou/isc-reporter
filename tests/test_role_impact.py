"""Role enablement impact."""

import csv
import json

import pytest
from openpyxl import load_workbook

from conftest import FakeApi, FakeIsc
from isc_reporter import cli, role_impact
from isc_reporter.errors import UsageError

AD, SAP, HR = "a" * 32, "s" * 32, "h" * 32
ROLE_ID = "r" * 32


def leaf(kind, attribute, operation, value, source=None):
    key = {"type": kind, "property": f"attribute.{attribute}"}
    if source:
        key["sourceId"] = source
    return {"operation": operation, "key": key, "stringValue": value}


def group(operation, *children):
    return {"operation": operation, "children": list(children)}


def entitlement(eid, name, source_id, source_name, privileged=False):
    return {"id": eid, "name": name, "type": "ENTITLEMENT", "attribute": "memberOf", "value": f"CN={name}",
            "source": {"id": source_id, "name": source_name}, "privileged": privileged}


USERS = entitlement("e1", "Finance Users", AD, "AD")
FI_DISPLAY = entitlement("e2", "FI_DISPLAY", SAP, "SAP", privileged=True)
SHARE = entitlement("e3", "Finance Share", AD, "AD")
ACCOUNTANTS = entitlement("e9", "Accountants", AD, "AD")
CRITERIA = group("AND",
                 group("OR", leaf("IDENTITY", "department", "EQUALS", "finance"),
                       leaf("ENTITLEMENT", "memberOf", "CONTAINS", "CN=Accountants", AD)),
                 leaf("IDENTITY", "cloudLifecycleState", "NOT_EQUALS", "inactive"))
ROLE = {"id": ROLE_ID, "name": "Finance", "enabled": False,
        "accessProfiles": [{"type": "ACCESS_PROFILE", "id": "ap1", "name": "Finance AD"},
                           {"type": "ACCESS_PROFILE", "id": "ap2", "name": "Finance SAP"}],
        "entitlements": [{"type": "ENTITLEMENT", "id": "e3", "name": "Finance Share"}],
        "membership": {"type": "STANDARD", "criteria": CRITERIA}}
ACCESS = [{**USERS, "via": ["Finance AD"]}, {**FI_DISPLAY, "via": ["Finance SAP"]},
          {**SHARE, "via": ["role"]}]


def person(name, department, access=(), accounts=("AD",), state="active"):
    sources = {"AD": AD, "SAP": SAP}
    return {"id": name, "name": name, "displayName": name.title(),
            "attributes": {"department": department, "cloudLifecycleState": state},
            "access": list(access),
            "accounts": [{"accountId": f"{name}-{s.lower()}", "source": {"id": sources[s], "name": s}}
                         for s in accounts]}


ALL = (USERS, FI_DISPLAY, SHARE)
IDENTITIES = [
    person("alice", "Finance"),                                      # gets the role and all of its access
    person("bob", "IT", (ACCOUNTANTS, *ALL)),                        # matches through an entitlement
    person("carol", "FINANCE", ALL),                                 # already has the role
    person("dan", "Finance", ALL, state="inactive"),                 # assigned by criteria, no longer matches
    person("eve", "Sales"),                                          # requested it
    person("frank", "Sales", ALL),                                   # has the access, not the role
    person("gina", "Sales", (USERS,)),                               # holds part of it: not listed
]
ASSIGNED = {"carol": "ROLE_MEMBERSHIP", "dan": "ROLE_MEMBERSHIP", "eve": "ACCESS_REQUEST"}


def data(**changes):
    values = {"role": ROLE, "access": ACCESS, "identities": IDENTITIES, "assigned": ASSIGNED}
    return role_impact.RoleData(**{**values, **changes})


def test_compare_ignores_case_and_orders_numbers():
    assert role_impact.compare("EQUALS", "Finance", "FINANCE")
    assert role_impact.compare("CONTAINS", "CN=Accountants,OU=Groups", "cn=accountants")
    assert role_impact.compare("STARTS_WITH", "FIN-01", "fin")
    assert role_impact.compare("ENDS_WITH", "FIN-01", "-01")
    assert role_impact.compare("GREATER_THAN", "10", "9")                 # as numbers, not as text
    assert role_impact.compare("LESS_THAN_EQUALS", "2026-01-01", "2026-06-30")
    assert not role_impact.compare("REGEX", "a", "a")


def test_criteria_on_identity_account_and_entitlement_values():
    alice = person("alice", ["Finance", "Audit"])
    alice["email"] = "alice@example.com"
    accounts = {HR: {"alice": [{"costCenter": "FIN-01"}]}}
    # one of the values matches; no value at all is "not equal"; a field outside "attributes"
    assert role_impact.matches(leaf("IDENTITY", "department", "EQUALS", "audit"), alice, {})
    assert not role_impact.matches(leaf("IDENTITY", "department", "NOT_EQUALS", "audit"), alice, {})
    assert role_impact.matches(leaf("IDENTITY", "title", "NOT_EQUALS", "boss"), alice, {})
    assert role_impact.matches(leaf("IDENTITY", "email", "ENDS_WITH", "@example.com"), alice, {})
    assert role_impact.matches(leaf("ACCOUNT", "costCenter", "STARTS_WITH", "fin", HR), alice, accounts)
    assert not role_impact.matches(leaf("ACCOUNT", "costCenter", "STARTS_WITH", "fin", AD), alice, accounts)
    bob = IDENTITIES[1]
    assert role_impact.matches(leaf("ENTITLEMENT", "memberOf", "EQUALS", "cn=accountants", AD), bob, {})
    assert not role_impact.matches(leaf("ENTITLEMENT", "memberOf", "EQUALS", "cn=accountants", SAP), bob, {})
    assert role_impact.members(ROLE, IDENTITIES, {}) == {"alice", "bob", "carol"}
    odd = {**ROLE, "membership": {"criteria": leaf("IDENTITY", "department", "REGEX", ".*")}}
    assert role_impact.unsupported(odd) == ["REGEX"] and role_impact.members(odd, IDENTITIES, {}) == set()


def test_identity_list_and_role_without_criteria():
    listed = {**ROLE, "membership": {"type": "IDENTITY_LIST", "identities": [{"id": "eve"}, {"id": "zoe"}]}}
    assert role_impact.members(listed, IDENTITIES, {}) == {"eve", "zoe"}
    result = role_impact.analyse(data(role={**ROLE, "membership": None}, assigned={}))
    assert result.matching == 0 and result.outcomes == {role_impact.GAP: 4}      # bob, carol, dan, frank
    assert dict(r.values() for r in role_impact.role_rows(result))["Membership"].startswith("none")


def test_outcomes_and_provisioning():
    result = role_impact.analyse(data())
    outcomes = {row["Identity"]: row["Outcome"] for row in result.identities}
    assert outcomes == {"alice": role_impact.NEW_ACCESS, "bob": role_impact.ROLE_ONLY,
                        "carol": role_impact.KEEPS, "dan": role_impact.LOSES, "eve": role_impact.OUTSIDE,
                        "frank": role_impact.GAP}
    assert [row["Identity"] for row in result.identities][:3] == ["alice", "bob", "dan"]   # worst first
    alice = result.identities[0]
    assert (alice["Entitlements to add"], alice["Accounts to create"], alice["Assigned by"]) == (3, 1, "")
    assert alice["To add"] == "AD: Finance Users\nSAP: FI_DISPLAY\nAD: Finance Share"
    assert {row["Identity"]: row["Assigned by"] for row in result.identities}["dan"] == "criteria"
    assert [(r["Source"], r["Entitlement"], r["Account"], r["Via"]) for r in result.provisioning] == [
        ("SAP", "FI_DISPLAY", role_impact.NEW_ACCOUNT, "Finance SAP"),               # privileged first
        ("AD", "Finance Share", "alice-ad", "role"),
        ("AD", "Finance Users", "alice-ad", "Finance AD")]
    assert (result.matching, result.new_members, result.new_accounts) == (3, 2, 1)
    users = next(row for row in result.entitlements if row["Entitlement"] == "Finance Users")
    counts = (users["To provision"], users["Held by matching identities"], users["Held by other identities"])
    assert counts == (1, 2, 3)                                  # alice / bob, carol / dan, frank, gina
    by_source = {row["Source"]: row for row in role_impact.by_source(result)}
    assert by_source["AD"]["Entitlements to add"] == 2 and by_source["AD"]["Accounts to create"] == 0
    assert (by_source["SAP"]["Accounts to create"], by_source["SAP"]["Privileged"]) == (1, 1)


def test_assignments_from_the_identities_access_when_not_read():
    carol = {**IDENTITIES[2], "access": [*ALL, {"type": "ROLE", "id": ROLE_ID, "name": "Finance"}]}
    result = role_impact.analyse(data(identities=[carol], assigned=None))
    assert [(r["Outcome"], r["Assigned by"]) for r in result.identities] == [(role_impact.KEEPS, "not known")]


def test_reports(tmp_path):
    result = role_impact.analyse(data())
    xlsx, pdf, csv_path = (str(tmp_path / f"r.{ext}") for ext in ("xlsx", "pdf", "csv"))
    role_impact.write_xlsx(xlsx, result)
    role_impact.write_pdf(pdf, result)
    role_impact.write_csv(csv_path, result)
    book = load_workbook(xlsx)
    assert book.sheetnames == ["Dashboard", "Identities", "Provisioning", "Entitlements", "By source", "Role",
                               "Chart data", "Read me"]
    assert book["Dashboard"]["B1"].value == "Role enablement impact: Finance"
    settings = {r[0]: r[1] for r in book["Role"].iter_rows(min_row=2, values_only=True)}
    assert settings["Membership"] == ('(identity.department EQUALS "finance" OR entitlement AD.memberOf '
                                      'CONTAINS "CN=Accountants") AND identity.cloudLifecycleState '
                                      'NOT_EQUALS "inactive"')
    assert settings["Access profiles"] == "Finance AD\nFinance SAP" and settings["Enabled"] == "no"
    assert open(pdf, "rb").read(4) == b"%PDF"
    with open(csv_path, encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.reader(fh, delimiter=";"))
    assert rows[0] == role_impact.IDENTITY_COLUMNS and len(rows) == 7


def test_load_role_from_files(tmp_path):
    (tmp_path / "role.json").write_text(json.dumps(ROLE), encoding="utf-8")
    assert role_impact.load_role([str(tmp_path / "role.json")])["name"] == "Finance"
    other = {**ROLE, "id": "x" * 32, "name": "Sales"}
    dimension = {**ROLE, "name": "Paris", "parentId": ROLE_ID}
    (tmp_path / "export.json").write_text(json.dumps({"objects": [{"object": other}, {"object": dimension}]}),
                                          encoding="utf-8")
    with pytest.raises(UsageError, match="2 roles"):
        role_impact.load_role([str(tmp_path)])
    assert role_impact.load_role([str(tmp_path)], "sales")["id"] == "x" * 32
    with pytest.raises(UsageError, match="no role named"):
        role_impact.load_role([str(tmp_path)], "Paris")


def fake_isc(role=ROLE):
    from sailpoint.roles.models import Role, RoleIdentity

    profiles = [{"id": "ap1", "name": "Finance AD", "source": {"id": AD, "name": "AD"},
                 "entitlements": [{k: v for k, v in USERS.items() if k not in ("source", "type")}]},
                {"id": "ap2", "name": "Finance SAP", "source": {"id": SAP, "name": "SAP"},
                 "entitlements": [{k: v for k, v in FI_DISPLAY.items() if k not in ("source", "type")}]}]
    documents = {"accessprofiles": profiles, "entitlements": [SHARE], "identities": IDENTITIES}
    assigned = [RoleIdentity.from_dict({"id": i, "name": i, "roleAssignmentSource": how})
                for i, how in ASSIGNED.items()]
    accounts = [{"id": "acc1", "identityId": "eve", "attributes": {"costCenter": "FIN-07"}}]
    return FakeIsc(
        roles=FakeApi(list_roles_v1=lambda **kw: [Role.from_dict(role)],
                      get_role_assigned_identities_v1=lambda id, limit, offset:
                      assigned if not offset else []),
        search=FakeApi(search_post_v1=lambda search, limit: documents[search.indices[0]]),
        accounts=FakeApi(list_accounts_v1=lambda filters, limit, offset: accounts if not offset else []))


def test_fetch_reads_access_assignments_and_accounts_of_account_criteria():
    with_account = {**ROLE, "membership": {"type": "STANDARD", "criteria": group(
        "OR", CRITERIA, leaf("ACCOUNT", "costCenter", "STARTS_WITH", "FIN", HR))}}
    isc = fake_isc(with_account)
    data = role_impact.fetch(isc, with_account, "attributes.department:*")
    assert [(e["name"], e["via"], e["source"]["name"]) for e in data.access] == [
        ("Finance Share", ["role"], "AD"), ("Finance Users", ["Finance AD"], "AD"),
        ("FI_DISPLAY", ["Finance SAP"], "SAP")]
    assert data.assigned == ASSIGNED and data.accounts == {HR: {"eve": [{"costCenter": "FIN-07"}]}}
    assert isc.accounts.called("list_accounts_v1")[0]["filters"] == f'sourceId eq "{HR}"'
    queries = [kw["search"].query.query for kw in isc.search.called("search_post_v1")]
    assert queries == ['id:("ap1" OR "ap2")', 'id:("e3")', "attributes.department:*"]
    outcomes = {r["Identity"]: r["Outcome"] for r in role_impact.analyse(data).identities}
    assert outcomes["eve"] == role_impact.KEEPS                   # now matches through her HR account


def test_cli(tmp_path, monkeypatch, capsys):
    monkeypatch.setattr("isc_reporter.client.connect", lambda config=None: fake_isc())
    out = tmp_path / "impact"
    args = ["role-impact", "Finance", "--format", "xlsx", "csv", "-o", str(out), "--log-dir", ""]
    assert cli.main(args) == 0
    rows = list(load_workbook(str(out) + ".xlsx")["Identities"].iter_rows(min_row=2, values_only=True))
    assert [r[0] for r in rows] == ["alice", "bob", "dan", "eve", "frank", "carol"]
    assert "enabling the role would assign it to 2, provision 3 entitlement(s) and create 1 account(s); " \
           "1 would lose it" in capsys.readouterr().err
    assert (tmp_path / "impact.csv").exists()


def test_cli_with_a_draft_role(tmp_path, monkeypatch):
    monkeypatch.setattr("isc_reporter.client.connect", lambda config=None: fake_isc())
    draft = {**ROLE, "membership": {"type": "STANDARD",
                                    "criteria": leaf("IDENTITY", "department", "EQUALS", "Sales")}}
    (tmp_path / "draft.json").write_text(json.dumps(draft), encoding="utf-8")
    base = ["role-impact", "--format", "csv", "--log-dir", "", "-o", str(tmp_path / "draft")]
    assert cli.main([*base, "--from", str(tmp_path / "draft.json")]) == 0
    with open(tmp_path / "draft.csv", encoding="utf-8-sig", newline="") as fh:
        outcomes = {row[0]: row[3] for row in list(csv.reader(fh, delimiter=";"))[1:]}
    assert outcomes["carol"] == outcomes["dan"] == role_impact.LOSES     # the new criteria leave them out
    assert outcomes["eve"] == role_impact.KEEPS                          # requested: she keeps it
    assert (outcomes["frank"], outcomes["gina"]) == (role_impact.ROLE_ONLY, role_impact.NEW_ACCESS)
    assert cli.main(base) == 2                                           # neither a name nor --from
