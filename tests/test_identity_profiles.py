"""Identity profile mapping workbook."""

import json

import pytest

from conftest import FakeApi, FakeIsc
from isc_reporter import cli, identity_profiles
from isc_reporter.errors import UsageError


def account(attribute, source="HR"):
    return {"type": "accountAttribute",
            "attributes": {"sourceName": source, "attributeName": attribute, "sourceId": "s" * 32}}


def mapping(name, definition):
    return {"identityAttributeName": name, "transformDefinition": definition}


PROFILE = {
    "id": "p" * 32, "name": "IP_HR", "priority": 20, "identityCount": 20,
    "owner": {"type": "IDENTITY", "id": "i" * 32, "name": "jdoe"},
    "authoritativeSource": {"type": "SOURCE", "id": "s" * 32, "name": "HR"},
    "identityAttributeConfig": {"enabled": True, "attributeTransforms": [
        mapping("uid", account("id")),
        mapping("endDate", {"type": "reference", "attributes": {
            "id": "Format date ISO8601", "input": account("endDate")}}),
        mapping("lifecycle", {"type": "reference", "attributes": {"id": "Lifecycle state"}}),
        mapping("email", {"type": "firstValid", "attributes": {"values": [
            account("mail", "AD"), account("e-mail"), {"type": "static", "attributes": {"value": "none"}}]}}),
        mapping("formula", {"type": "static", "attributes": {"value": "=1+1"}}),
    ]},
}
TRANSFORMS = {
    "Format date ISO8601": {"name": "Format date ISO8601", "type": "dateFormat",
                            "attributes": {"inputFormat": "dd/MM/yyyy", "outputFormat": "ISO8601"}},
    "Lifecycle state": {"name": "Lifecycle state", "type": "dateCompare",
                        "attributes": {"firstDate": account("startDate"), "operator": "LTE"}},
}


def test_mapping_rows_explain_each_attribute():
    rows = {r["identity_attribute"]: r for r in identity_profiles.mapping_rows(PROFILE, TRANSFORMS)}
    uid = rows["uid"]
    assert (uid["mapping"], uid["display_name"], uid["sources"], uid["source_attributes"]) == \
        ("Direct copy", "Username", "HR", "id")
    assert uid["explanation"] == "Copied from the account attribute 'id' of the source HR."
    end = rows["endDate"]
    assert (end["mapping"], end["transforms"], end["logic"]) == \
        ("Named transform", "Format date ISO8601", "Format date ISO8601(HR.endDate)")
    assert end["explanation"] == ("Value of 'endDate' from HR, converted by the transform "
                                  "'Format date ISO8601' (changes a date format).")
    # a referenced transform without input: its own account attributes are found
    assert rows["lifecycle"]["source_attributes"] == "startDate"
    assert rows["lifecycle"]["explanation"].endswith("which reads 'startDate' from HR.")
    email = rows["email"]
    assert email["mapping"] == "Inline transform"
    assert email["logic"] == 'firstValid(AD.mail, HR.e-mail, "none")'
    assert email["explanation"] == "Takes the first non-empty value, using 'mail' from AD, 'e-mail' from HR."
    assert email["source_attributes"] == "AD.mail\nHR.e-mail"
    assert "Other source" in email["notes"]
    assert rows["formula"]["explanation"] == 'Always "=1+1".'


def test_checks_flag_likely_mistakes():
    profile = {**PROFILE, "identityRefreshRequired": True,
               "identityAttributeConfig": {"attributeTransforms": [
                   mapping("email", account("e-mail")),
                   mapping("cloudLifecycleState", account("e-mail")),
                   mapping("x", {"type": "reference", "attributes": {"id": "Gone"}}),
               ]}}
    rows = identity_profiles.mapping_rows(profile, {"Other": {}})
    checks = {(c[1], c[2]): c[0] for c in identity_profiles.profile_checks(profile, rows)}
    assert checks[("uid", "Required attribute not mapped")] == "Warning"
    assert checks[("manager", "Common attribute not mapped")] == "Info"
    assert checks[("cloudLifecycleState", "Lifecycle state copied as is")] == "Warning"
    assert checks[("email", "Same mapping")] == "Info"
    assert checks[("x", "Missing transform")] == "Warning"
    assert checks[("", "Identity refresh pending")] == "Info"
    # without any transform definitions, a reference cannot be called missing
    assert "Missing transform" not in identity_profiles.mapping_rows(profile, {})[2]["notes"]


def test_load_files_accepts_single_lists_and_config_exports(tmp_path):
    (tmp_path / "profile.json").write_text(json.dumps(PROFILE), encoding="utf-8")
    (tmp_path / "export.json").write_text(json.dumps({"version": 1, "objects": [
        {"self": {"type": "IDENTITY_PROFILE"}, "object": {**PROFILE, "name": "IP_Other"}},
        {"self": {"type": "TRANSFORM"}, "object": TRANSFORMS["Lifecycle state"]}]}), encoding="utf-8")
    (tmp_path / "transforms.json").write_text(json.dumps([TRANSFORMS["Format date ISO8601"]]))
    profiles, transforms = identity_profiles.load_files([str(tmp_path)])
    assert sorted(p["name"] for p in profiles) == ["IP_HR", "IP_Other"]
    assert sorted(transforms) == ["Format date ISO8601", "Lifecycle state"]
    with pytest.raises(UsageError):
        identity_profiles.load_files([str(tmp_path / "transforms.json")])


def test_workbook_sheets_and_summary(tmp_path):
    import zipfile

    from openpyxl import load_workbook

    path = str(tmp_path / "map.xlsx")
    other = {**PROFILE, "name": "IP/HR:test", "priority": 10}
    assert identity_profiles.write_xlsx(path, [PROFILE, other], TRANSFORMS, {"uid": "UID"},
                                        origin="tenant") == 2
    book = load_workbook(path)
    assert book.sheetnames == ["Dashboard", "Profiles", "IP_HR_test", "IP_HR", "Matrix", "Coverage",
                               "Attribute usage", "Usage detail", "Transforms", "Transform dependencies",
                               "Checks", "Chart data", "Read me"]
    sheet = book["IP_HR"]
    assert [c.value for c in sheet[1]][:4] == ["Identity attribute", "Display name", "Mapping", "Explanation"]
    assert [c.value for c in sheet[2]][:3] == ["uid", "UID", "Direct copy"]
    formula = next(row for row in sheet.iter_rows(min_row=2) if row[0].value == "formula")
    assert formula[7].value == '"=1+1"' and formula[7].data_type == "s"   # text, never a formula
    profiles = {row[0]: row for row in book["Profiles"].iter_rows(min_row=2, values_only=True)}
    assert profiles["IP_HR"][5:8] == (5, 1, 2)          # attributes, direct copy, named transform
    assert [c.value for c in book["Transforms"]["A"]][1:] == ["Format date ISO8601", "Lifecycle state"]
    assert {c.value for c in book["Checks"]["D"]} >= {"Common attribute not mapped", "Other source"}
    package = zipfile.ZipFile(path)
    assert sum(n.startswith("xl/charts/chart") for n in package.namelist()) == 4
    dashboard = load_workbook(path)["Dashboard"]
    assert dashboard["B4"].value == "Identity profiles" and dashboard["B5"].value == "2"


POLICIES = [{"source": "AD", "usageType": "CREATE", "fields": [
    {"name": "mail", "transform": {"type": "identityAttribute", "attributes": {"name": "email"}}},
    {"name": "cn", "transform": {"type": "reference", "attributes": {"id": "Full name"}}},
    {"name": "free", "type": "string"}]}]
TRANSFORMS_WITH_READERS = {**TRANSFORMS, "Full name": {
    "name": "Full name", "type": "concat", "attributes": {"values": [
        {"type": "identityAttribute", "attributes": {"name": "firstname"}},
        {"type": "static", "attributes": {"value": " $lastname"}}]}}}


def test_attribute_usage_across_profiles_transforms_and_policies(tmp_path):
    from isc_reporter import attribute_usage

    profile = {**PROFILE, "identityAttributeConfig": {"attributeTransforms": [
        mapping("email", account("mail")),
        mapping("firstname", account("first")),
        mapping("lastname", account("last")),
        mapping("displayName", {"type": "identityAttribute", "attributes": {"name": "firstname"}})]}}
    rows = [identity_profiles.mapping_rows(profile, TRANSFORMS_WITH_READERS)]
    known = {"email", "firstname", "lastname", "displayName"}
    usages = attribute_usage.collect([profile], rows, TRANSFORMS_WITH_READERS, POLICIES, known)
    reads = {(u["attribute"], u["where"], u["object"], u["location"]) for u in usages if u["role"] == "Read"}
    assert reads == {
        ("firstname", "Identity profile", "IP_HR", "mapping of 'displayName'"),
        ("firstname", "Transform", "Full name", "type concat"),
        ("lastname", "Transform", "Full name", "type concat"),
        ("email", "Provisioning policy", "AD / CREATE", "field 'mail'"),
        ("firstname", "Provisioning policy", "AD / CREATE", "field 'cn'"),   # through the named transform
        ("lastname", "Provisioning policy", "AD / CREATE", "field 'cn'"),
    }
    summary = {r["Identity attribute"]: r for r in attribute_usage.summary_rows(
        usages, ["IP_HR"], lambda a: "")}
    assert summary["email"]["Mapped in"] == "1 / 1: IP_HR"
    assert summary["email"]["Used in provisioning policies"] == "AD / CREATE: field 'mail'"
    assert summary["email"]["Sources"] == "AD"
    assert summary["firstname"]["Read by transforms"] == "Full name"
    assert summary["firstname"]["Uses"] == 3

    from openpyxl import load_workbook

    path = str(tmp_path / "usage.xlsx")
    identity_profiles.write_xlsx(path, [profile], TRANSFORMS_WITH_READERS, policies=POLICIES)
    book = load_workbook(path)
    sheet = book["IP_HR"]
    header = [c.value for c in sheet[1]]
    used_in = header.index("Also used in")
    email = next(r for r in sheet.iter_rows(min_row=2, values_only=True) if r[0] == "email")
    assert "Provisioning policy AD / CREATE: field 'mail'" in email[used_in]
    assert [c.value for c in book["Usage detail"][1]] == [
        "Identity attribute", "Where", "Object", "Location", "Role", "Detail"]


def test_load_policies_from_source_export(tmp_path):
    (tmp_path / "source.json").write_text(json.dumps({"objects": [{"object": {
        "name": "AD", "provisioningPolicies": POLICIES[0:1][0:1] and [
            {"usageType": "CREATE", "fields": POLICIES[0]["fields"]}]}}]}), encoding="utf-8")
    policies = identity_profiles.load_policies([str(tmp_path)])
    assert [(p["source"], p["usageType"]) for p in policies] == [("AD", "CREATE")]


def test_pdf(tmp_path):
    path = str(tmp_path / "map.pdf")
    assert identity_profiles.write_pdf(path, [PROFILE], TRANSFORMS) == 1
    assert open(path, "rb").read(4) == b"%PDF"


def test_cli_from_tenant(tmp_path, monkeypatch, capsys):
    attribute = type("A", (), {"name": "uid", "display_name": "User ID"})()
    isc = FakeIsc(identity_profiles=FakeApi(list_identity_profiles_v1=lambda **kw: []),
                  transforms=FakeApi(list_transforms_v1=lambda **kw: []),
                  identity_attributes=FakeApi(list_identity_attributes_v1=lambda **kw: [attribute]))
    from sailpoint.identity_profiles.models import IdentityProfile

    model = IdentityProfile.from_dict(PROFILE)
    isc.identity_profiles.responses["list_identity_profiles_v1"] = \
        lambda limit, offset: [model] if offset == 0 else []
    monkeypatch.setattr("isc_reporter.client.connect", lambda config=None: isc)
    out = tmp_path / "m.xlsx"
    assert cli.main(["mappings", "-o", str(out), "--log-dir", ""]) == 0
    assert out.exists()
    # the old names still work, with a warning (an unknown profile: exit code 2)
    for old in ("identity-profiles", "profiles"):
        assert cli.main([old, "--profile", "nope", "-o", str(out), "--log-dir", ""]) == 2
        assert f"The '{old}' command is now called 'mappings'" in capsys.readouterr().err
