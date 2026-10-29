"""Dependencies between named transforms."""

from openpyxl import load_workbook

from isc_reporter import identity_profiles, transform_graph, transform_tree


def account(attribute, source="HR"):
    return {"type": "accountAttribute", "attributes": {"sourceName": source, "attributeName": attribute}}


def reference(name, given=None):
    return {"type": "reference", "attributes": {"id": name, **({"input": given} if given else {})}}


def identity(name):
    return {"type": "identityAttribute", "attributes": {"name": name}}


def profile(name, priority, mappings):
    return {"name": name, "priority": priority, "identityAttributeConfig": {"attributeTransforms": [
        {"identityAttributeName": attribute, "transformDefinition": definition}
        for attribute, definition in mappings]}}


TRANSFORMS = {
    "Accents": {"name": "Accents", "type": "decomposeDiacriticalMarks", "attributes": {}},
    "Name": {"name": "Name", "type": "normalizeNames", "attributes": {"input": reference("Accents")}},
    "Email": {"name": "Email", "type": "lower", "attributes": {"input": {"type": "concat", "attributes": {
        "values": [identity("firstname"), "@example.com"]}}}},
    "Greeting": {"name": "Greeting", "type": "static", "attributes": {"value": "Hello $firstname"}},
    "Loop A": {"name": "Loop A", "type": "upper", "attributes": {"input": reference("Loop B")}},
    "Loop B": {"name": "Loop B", "type": "lower", "attributes": {"input": reference("Loop A")}},
}
PROFILES = [
    profile("Employees", 10, [("firstname", reference("Name", account("first"))),
                              ("email", reference("Email")), ("uid", account("id")),
                              ("title", reference("Gone"))]),
    profile("Contractors", 20, [("firstname", account("first", "CR")),
                                ("lastname", reference("Name", account("last", "CR")))]),
]


def graph():
    profiles, rows, _ = identity_profiles.analyse(PROFILES, TRANSFORMS)
    return transform_graph.build(profiles, rows, TRANSFORMS), profiles, rows


def test_edges_between_transforms_and_attributes():
    found, _, _ = graph()
    edges = {(e["from"], e["to"], e["kind"]) for e in found.edges}
    assert ("t:Accents", "t:Name", "calls") in edges
    assert {("t:Name", "a:Employees/firstname", "uses"),
            ("t:Name", "a:Contractors/lastname", "uses")} <= edges
    # Email is used by Employees only: it reads the firstname of Employees, not of Contractors
    assert ("a:Employees/firstname", "t:Email", "reads") in edges
    assert ("a:Contractors/firstname", "t:Email", "reads") not in edges
    # a transform used nowhere reads the attribute in every profile (here through Velocity)
    assert {("a:Employees/firstname", "t:Greeting", "reads"),
            ("a:Contractors/firstname", "t:Greeting", "reads")} <= edges
    assert found.nodes["t:Gone"]["kind"] == transform_graph.MISSING
    assert found.nodes["t:Name"]["what"] == "normalizeNames: normalizes person names"
    assert "a:Employees/uid" not in found.nodes                     # no named transform involved
    assert transform_graph.used_in(found, "t:Accents") == {"Employees", "Contractors"}


def test_trees_follow_dependencies_across_attributes():
    found, _, _ = graph()
    email = transform_tree.draw(transform_graph.tree(found, "a:Employees/email")).splitlines()
    assert email == [
        "email",
        "└─ uses: Email",
        "   └─ reads: firstname (Employees)",
        "      └─ uses: Name",
        "         └─ calls: Accents"]
    loop = transform_graph.tree(found, "t:Loop A")
    assert loop.children[0].children[0].what == "already shown above: circular dependency"


def test_roots_list_attributes_then_unused_transforms_even_in_a_loop():
    found, profiles, rows = graph()
    assert transform_graph.roots(found, profiles, rows) == [
        ("Employees: firstname", "a:Employees/firstname"), ("Employees: email", "a:Employees/email"),
        ("Employees: title", "a:Employees/title"), ("Contractors: lastname", "a:Contractors/lastname"),
        (transform_graph.NOT_USED, "t:Greeting"), (transform_graph.NOT_USED, "t:Loop A")]


def test_workbook_and_pdf(tmp_path):
    path = str(tmp_path / "mappings.xlsx")
    identity_profiles.write_xlsx(path, PROFILES, TRANSFORMS)
    book = load_workbook(path)
    sheet = book["Transform dependencies"]
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0] == ("Used by", "Tree", "What it is", "Level")
    assert [r[1] for r in rows[1:4]] == ["firstname", "└─ uses: Name", "   └─ calls: Accents"]
    assert [sheet.row_dimensions[r].outlineLevel for r in range(2, 5)] == [0, 1, 2]
    header = [c.value for c in book["Transforms"][1]]
    transforms = {r[0]: dict(zip(header, r, strict=True))
                  for r in book["Transforms"].iter_rows(min_row=2, values_only=True)}
    assert transforms["Name"]["Depends on"].splitlines() == ["Name", "└─ calls: Accents"]
    assert transforms["Accents"]["Called by"] == "Name"
    assert transforms["Loop A"]["Uses"] == 0                     # not used, listed all the same
    assert identity_profiles.write_pdf(str(tmp_path / "mappings.pdf"), PROFILES, TRANSFORMS) == 2
